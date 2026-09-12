from flask import request, send_file
from admin.modelo.pagos import PagosModelo
from admin.modelo.empresa import EmpresaModelo
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Drawing, Line
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
import tempfile
import os


def _logo_path():
    return os.path.abspath(os.path.join(
        os.path.dirname(os.path.abspath(__file__)), '..', '..', 'static', 'Logo SAC.jpeg'
    ))


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 9)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 12 * mm, 'Sistema de Gestión | Reporte generado automáticamente')
    canvas.drawRightString(195 * mm, 12 * mm, f'Página {doc.page}')
    canvas.restoreState()


def borde(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.lightgrey)
    canvas.setLineWidth(1)
    canvas.rect(20 * mm, 20 * mm, LETTER[0] - 40 * mm, LETTER[1] - 40 * mm)
    canvas.restoreState()


def _build_filters(desde, hasta, cod_tipo_pago, status):
    condiciones = []
    parametros = []

    if desde:
        condiciones.append('p.fecha_pago >= %s')
        parametros.append(desde)

    if hasta:
        condiciones.append('p.fecha_pago <= %s')
        parametros.append(hasta)

    if cod_tipo_pago and cod_tipo_pago > 0:
        condiciones.append('dp.cod_tipo_pago = %s')
        parametros.append(cod_tipo_pago)

    if status == 'active':
        condiciones.append('p.status = 1')
    elif status == 'anulado':
        condiciones.append('p.status = 0')

    where = ' AND '.join(condiciones) if condiciones else '1=1'
    return where, parametros


def _fetch_rows(modelo, sql, parametros):
    cursor = modelo.conexion.cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    filas = cursor.fetchall()
    cursor.close()
    return filas


def _consultar_pagos(modelo, desde=None, hasta=None, cod_tipo_pago=0, status='all'):
    where, parametros = _build_filters(desde, hasta, cod_tipo_pago, status)
    sql = f'''
        SELECT
            p.cod_pago,
            p.fecha_pago,
            p.monto,
            p.status,
            u.cedula,
            u.nombre_usuario AS nombre_cliente,
            tp.nombre_tipo_pago AS nombre_tipo_pago,
            m.nombre_moneda AS nombre_moneda
        FROM pago p
        INNER JOIN servicio s ON p.cod_servicio = s.cod_servicio
        INNER JOIN seguridad.usuario u ON s.cedula = u.cedula
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN tipo_pago tp ON tp.cod_tipo_pago = dp.cod_tipo_pago
        LEFT JOIN moneda m ON m.cod_moneda = COALESCE(dp.cod_moneda, 1)
        WHERE {where}
        ORDER BY p.fecha_pago DESC
    '''
    return _fetch_rows(modelo, sql, parametros)


def _generar_pdf(empresa, listado, filtros):
    archivo = os.path.join(tempfile.gettempdir(), f'reporte_pagos_{int(datetime.now().timestamp())}.pdf')
    doc = SimpleDocTemplate(
        archivo,
        pagesize=LETTER,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm
    )
    styles = getSampleStyleSheet()

    empresa_style = ParagraphStyle(
        'empresa', parent=styles['Heading1'], fontName='Helvetica-Bold',
        fontSize=18, leading=20, textColor=colors.HexColor('#1B365D'), alignment=TA_LEFT
    )
    titulo_style = ParagraphStyle(
        'titulo', parent=styles['Heading1'], fontName='Helvetica-Bold',
        fontSize=18, textColor=colors.HexColor('#B8860B'), alignment=TA_CENTER, spaceAfter=4
    )
    subtitulo_style = ParagraphStyle(
        'subtitulo', parent=styles['Normal'], fontName='Helvetica', fontSize=10,
        textColor=colors.grey, alignment=TA_CENTER
    )

    normal_style = ParagraphStyle(
        "normal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14
    )

    elementos = []

    logo = ''
    logo_path = _logo_path()
    if os.path.exists(logo_path):
        logo = Image(logo_path, width=42 * mm, height=42 * mm)

    direccion_empresa = empresa.get('direccion') or 'Carrera 13 entre calles 55 y 56'
    telefono_empresa = empresa.get('telefono') or '04145513219'
    correo_empresa = empresa.get('email') or 'guilermina@gmail.com'

    empresa_html = f"""
        <font size="18" color="#1B365D"><b>ESTUDIO DE BELLEZA GUILLERMINA</b></font><br/><br/>
        <font size="10"><b>Dirección:</b> {direccion_empresa}<br/>
        <b>Teléfono:</b> {telefono_empresa}<br/>
        <b>Correo:</b> {correo_empresa}</font>
    """
    encabezado = Table([[logo, Paragraph(empresa_html, normal_style)]], colWidths=[48 * mm, 120 * mm])
    encabezado.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8)
    ]))
    elementos.append(encabezado)

    linea = Drawing(520, 12)
    linea.add(Line(0, 5, 520, 5, strokeColor=colors.HexColor('#D4AF37'), strokeWidth=2))
    elementos.extend([linea, Spacer(1, 12)])
    elementos.append(Paragraph('REPORTE GENERAL DE PAGOS', titulo_style))
    elementos.append(Paragraph('Listado de pagos registrados en el sistema administrativo.', subtitulo_style))
    elementos.append(Paragraph(f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y')}", subtitulo_style))
    elementos.append(Spacer(1, 10))

    data = [['Cédula', 'Cliente', 'Fecha', 'Monto', 'Tipo', 'Moneda', 'Status']]
    for pago in listado:
        status_lbl = 'Activo' if pago['status'] == 1 else 'Anulado'
        data.append([
            pago['cedula'],
            pago['nombre_cliente'],
            pago['fecha_pago'],
            f"{pago['monto']:.2f}",
            pago['nombre_tipo_pago'] or 'N/A',
            pago['nombre_moneda'] or 'Bs',
            status_lbl,
        ])

    tabla = Table(
        data,
        repeatRows=1,
        hAlign='LEFT',
        colWidths=[26 * mm, 30 * mm, 23 * mm, 18 * mm, 26 * mm, 18 * mm, 12 * mm]
    )
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
        ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 1),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
    ]))
    elementos.append(tabla)

    doc.build(
        elementos,
        onFirstPage=lambda c, d: (borde(c, d), footer(c, d)),
        onLaterPages=lambda c, d: (borde(c, d), footer(c, d))
    )
    return archivo


def _generar_excel(empresa, listado, filtros):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Pagos'

    logo_path = _logo_path()
    if os.path.exists(logo_path):
        logo = ExcelImage(logo_path)
        logo.width = 75
        logo.height = 60
        sheet.add_image(logo, 'D1')

    sheet.merge_cells('A6:G6')
    sheet['A6'] = 'ESTUDIO DE BELLEZA GUILLERMINA'
    sheet['A6'].font = Font(bold=True, size=14)
    sheet['A6'].alignment = Alignment(horizontal='center')
    for row, value in enumerate([
        empresa.get('direccion', ''),
        empresa.get('telefono', ''),
        empresa.get('email', ''),
        datetime.now().strftime('%d-%m-%Y'),
    ], start=7):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        sheet.cell(row=row, column=1).value = value
        sheet.cell(row=row, column=1).alignment = Alignment(horizontal='center')

    sheet.row_dimensions[1].height = 50
    sheet.row_dimensions[2].height = 8
    sheet.row_dimensions[3].height = 8
    sheet.row_dimensions[4].height = 8
    sheet.row_dimensions[5].height = 8
    sheet.row_dimensions[6].height = 22
    gold_border = Border(bottom=Side(style='medium', color='D4AF37'))
    for cell in sheet[10]:
        cell.border = gold_border
    sheet['A10'].font = Font(italic=True, color='777777')

    sheet.append([])
    sheet.append(['Reporte de Pagos'])
    sheet.merge_cells('A12:G12')
    sheet['A12'].font = Font(bold=True, size=14, color='B8860B')
    sheet['A12'].alignment = Alignment(horizontal='center')
    sheet.append(['Emitido', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    sheet.append([])
    sheet.append(['Cédula', 'Cliente', 'Fecha', 'Monto', 'Tipo de Pago', 'Moneda', 'Status'])

    for pago in listado:
        sheet.append([
            pago['cedula'],
            pago['nombre_cliente'],
            pago['fecha_pago'],
            float(pago['monto']),
            pago['nombre_tipo_pago'] or 'N/A',
            pago['nombre_moneda'] or 'Bs',
            'Activo' if pago['status'] == 1 else 'Anulado',
        ])

    for col in range(1, 8):
        cell = sheet.cell(row=15, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1B365D')
        cell.alignment = Alignment(horizontal='center')

    sheet.column_dimensions['A'].width = 16
    sheet.column_dimensions['B'].width = 24
    sheet.column_dimensions['C'].width = 16
    sheet.column_dimensions['D'].width = 12
    sheet.column_dimensions['E'].width = 18
    sheet.column_dimensions['F'].width = 14
    sheet.column_dimensions['G'].width = 12

    archivo = os.path.join(tempfile.gettempdir(), f'reporte_pagos_{int(datetime.now().timestamp())}.xlsx')
    workbook.save(archivo)
    return archivo


def handle_request(ruta):
    modelo = PagosModelo()

    desde = request.form.get('desde') or request.args.get('desde')
    hasta = request.form.get('hasta') or request.args.get('hasta')
    cod_tipo_pago = int(request.form.get('cod_tipo_pago', 0) or 0)
    status = request.form.get('status', 'all')

    filtros = {
        'desde': desde or '',
        'hasta': hasta or '',
        'cod_tipo_pago': cod_tipo_pago,
        'status': status,
    }

    registro = _consultar_pagos(modelo, desde=desde, hasta=hasta, cod_tipo_pago=cod_tipo_pago, status=status)
    tipopagos = modelo.listar_tipos_pago()

    if request.method == 'POST' and request.form.get('pdf'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_pdf(empresa_data, registro, filtros)
        return send_file(archivo, mimetype='application/pdf', as_attachment=False)

    if request.method == 'POST' and request.form.get('excel'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_excel(empresa_data, registro, filtros)
        return send_file(archivo, as_attachment=True, download_name='reporte_pagos.xlsx')

    return {
        'registro': registro,
        'tipopagos': tipopagos,
        'filtros': filtros,
    }
