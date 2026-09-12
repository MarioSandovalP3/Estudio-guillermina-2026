from flask import request, send_file
from admin.modelo.producto import ProductoModelo
from admin.modelo.empresa import EmpresaModelo
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
import tempfile
import os


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 9)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 12 * mm, 'Sistema de Gestión | Reporte generado automáticamente')
    canvas.drawRightString(195 * mm, 12 * mm, f'Página {doc.page}')
    canvas.restoreState()


def borde(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.lightgrey)
    canvas.setLineWidth(1)
    canvas.rect(20 * mm, 20 * mm, LETTER[0] - 40 * mm, LETTER[1] - 40 * mm)
    canvas.restoreState()


def _build_filters(cod_categoria, cod_marca, status):
    condiciones = []
    parametros = []

    if cod_categoria and cod_categoria > 0:
        condiciones.append('p.cod_categoria = %s')
        parametros.append(cod_categoria)

    if cod_marca and cod_marca > 0:
        condiciones.append('p.cod_marca = %s')
        parametros.append(cod_marca)

    if status == 'active':
        condiciones.append('p.status = 1')
    elif status == 'inactive':
        condiciones.append('p.status = 0')

    where = ' AND '.join(condiciones) if condiciones else '1=1'
    return where, parametros


def _fetch_scalar(modelo, sql, parametros):
    cursor = modelo.conn().cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    fila = cursor.fetchone()
    cursor.close()
    return fila


def _fetch_rows(modelo, sql, parametros):
    cursor = modelo.conn().cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    filas = cursor.fetchall()
    cursor.close()
    return filas


def _consultar_estadisticas(modelo, cod_categoria=0, cod_marca=0, status='all'):
    where, parametros = _build_filters(cod_categoria, cod_marca, status)

    resumen = _fetch_scalar(modelo, f'''
        SELECT
            COUNT(*) AS total_productos,
            COALESCE(SUM(stock), 0) AS stock_total,
            COALESCE(SUM(precio_producto * stock), 0) AS valor_total,
            SUM(stock <= 5) AS stock_bajo
        FROM producto p
        WHERE {where}
    ''', parametros) or {}

    por_categoria = _fetch_rows(modelo, f'''
        SELECT
            c.cod_categoria,
            c.nombre_categoria,
            COUNT(*) AS total_productos,
            COALESCE(SUM(p.stock), 0) AS stock_total
        FROM producto p
        INNER JOIN categoria c ON p.cod_categoria = c.cod_categoria
        WHERE {where}
        GROUP BY c.cod_categoria, c.nombre_categoria
        ORDER BY stock_total DESC
    ''', parametros)

    por_marca = _fetch_rows(modelo, f'''
        SELECT
            m.cod_marca,
            m.nombre_marca,
            COUNT(*) AS total_productos,
            COALESCE(SUM(p.stock), 0) AS stock_total
        FROM producto p
        INNER JOIN marca m ON p.cod_marca = m.cod_marca
        WHERE {where}
        GROUP BY m.cod_marca, m.nombre_marca
        ORDER BY stock_total DESC
    ''', parametros)

    stock_bajo = _fetch_rows(modelo, f'''
        SELECT
            p.cod_producto,
            p.nombre_producto,
            p.stock,
            p.precio_producto
        FROM producto p
        WHERE {where} AND p.stock <= 5
        ORDER BY p.stock ASC, p.cod_producto DESC
        LIMIT 10
    ''', parametros)

    ultimos_productos = _fetch_rows(modelo, f'''
        SELECT
            p.cod_producto,
            p.nombre_producto,
            p.stock,
            p.precio_producto,
            p.status,
            c.nombre_categoria,
            m.nombre_marca
        FROM producto p
        INNER JOIN categoria c ON p.cod_categoria = c.cod_categoria
        INNER JOIN marca m ON p.cod_marca = m.cod_marca
        WHERE {where}
        ORDER BY p.cod_producto DESC
        LIMIT 10
    ''', parametros)

    return {
        'total_productos': int(resumen.get('total_productos', 0)),
        'stock_total': float(resumen.get('stock_total', 0.0)),
        'valor_total': float(resumen.get('valor_total', 0.0)),
        'stock_bajo': int(resumen.get('stock_bajo', 0)),
        'por_categoria': por_categoria,
        'por_marca': por_marca,
        'stock_bajo_items': stock_bajo,
        'ultimos_productos': ultimos_productos,
    }


def _listar_catalogos(modelo):
    cursor = modelo.conn().cursor(dictionary=True)
    cursor.execute('SELECT cod_categoria, nombre_categoria FROM categoria ORDER BY nombre_categoria')
    categorias = cursor.fetchall()
    cursor.execute('SELECT cod_marca, nombre_marca FROM marca ORDER BY nombre_marca')
    marcas = cursor.fetchall()
    cursor.close()
    return categorias, marcas


def _generar_pdf(empresa, estadisticas, filtros):
    archivo = os.path.join(tempfile.gettempdir(), f'reporte_productos_estadistico_{int(datetime.now().timestamp())}.pdf')
    doc = SimpleDocTemplate(
        archivo,
        pagesize=LETTER,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm
    )
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle('titulo', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=18, textColor=colors.HexColor('#1B365D'), alignment=1, spaceAfter=6)
    normal_style = ParagraphStyle('normal', parent=styles['Normal'], fontName='Helvetica', fontSize=10, leading=14)

    elementos = []

    if empresa.get('logo'):
        base_dir = os.path.dirname(os.path.abspath(__file__))
        logo_path = os.path.abspath(os.path.join(base_dir, '..', '..', 'static', 'assets', 'img', empresa['logo']))
        if os.path.exists(logo_path):
            imagen = Image(logo_path, width=50 * mm, height=30 * mm)
            imagen.hAlign = 'LEFT'
            elementos.append(imagen)
            elementos.append(Spacer(1, 8))

    elementos.append(Paragraph('Estudio MT - Salón & Estética', titulo_style))
    elementos.append(Paragraph(f"Dirección: {empresa.get('direccion','')}" if empresa.get('direccion') else 'Dirección: N/A', normal_style))
    elementos.append(Paragraph(f"Teléfono: {empresa.get('telefono','')}" if empresa.get('telefono') else 'Teléfono: N/A', normal_style))
    elementos.append(Paragraph(f"Email: {empresa.get('email','')}" if empresa.get('email') else 'Email: N/A', normal_style))
    elementos.append(Spacer(1, 12))

    elementos.append(Paragraph('<b>REPORTE ESTADÍSTICO DE PRODUCTOS</b>', styles['Heading2']))
    elementos.append(Paragraph(f"Fecha: {datetime.now().strftime('%d-%m-%Y')}", normal_style))
    elementos.append(Spacer(1, 12))

    filtros_data = [
        ['Categoría', filtros.get('cod_categoria', 'Todos')],
        ['Marca', filtros.get('cod_marca', 'Todos')],
        ['Estado', filtros.get('status', 'Todos')],
    ]
    filtros_table = Table(filtros_data, hAlign='LEFT', colWidths=[30 * mm, 115 * mm])
    filtros_table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
    ]))
    elementos.extend([filtros_table, Spacer(1, 12)])

    metrics_data = [
        ['Métrica', 'Valor'],
        ['Total de productos', str(estadisticas['total_productos'])],
        ['Stock total', f"{estadisticas['stock_total']:.2f}"],
        ['Valor total', f"{estadisticas['valor_total']:.2f}"],
        ['Productos con stock bajo', str(estadisticas['stock_bajo'])],
    ]
    metrics_table = Table(metrics_data, hAlign='LEFT', colWidths=[80 * mm, 65 * mm])
    metrics_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
    ]))
    elementos.extend([metrics_table, Spacer(1, 12)])

    if estadisticas['por_categoria']:
        elementos.append(Paragraph('Por categoría', styles['Heading3']))
        categoria_data = [['Categoría', 'Productos', 'Stock']]
        for item in estadisticas['por_categoria']:
            categoria_data.append([item['nombre_categoria'], item['total_productos'], f"{item['stock_total']:.2f}"])
        categoria_table = Table(categoria_data, hAlign='LEFT', colWidths=[75 * mm, 35 * mm, 45 * mm], repeatRows=1)
        categoria_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
        ]))
        elementos.extend([categoria_table, Spacer(1, 12)])

    if estadisticas['por_marca']:
        elementos.append(Paragraph('Por marca', styles['Heading3']))
        marca_data = [['Marca', 'Productos', 'Stock']]
        for item in estadisticas['por_marca']:
            marca_data.append([item['nombre_marca'], item['total_productos'], f"{item['stock_total']:.2f}"])
        marca_table = Table(marca_data, hAlign='LEFT', colWidths=[75 * mm, 35 * mm, 45 * mm], repeatRows=1)
        marca_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
        ]))
        elementos.extend([marca_table, Spacer(1, 12)])

    doc.build(
        elementos,
        onFirstPage=lambda c, d: (borde(c, d), footer(c, d)),
        onLaterPages=lambda c, d: (borde(c, d), footer(c, d))
    )
    return archivo


def _generar_excel(empresa, estadisticas, filtros):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Estadísticas'

    sheet.merge_cells('A1:B1')
    sheet['A1'] = empresa.get('nombre', 'Estudio MT')
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A1'].alignment = Alignment(horizontal='center')

    sheet['A2'] = empresa.get('direccion', '')
    sheet['A3'] = empresa.get('telefono', '')
    sheet['A4'] = empresa.get('email', '')
    sheet['A5'] = datetime.now().strftime('%d-%m-%Y')

    sheet.append([])
    sheet.append(['Reporte Estadístico de Productos'])
    sheet.append(['Emitido', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    sheet.append([])
    sheet.append(['Categoría', filtros.get('cod_categoria', 'Todos')])
    sheet.append(['Marca', filtros.get('cod_marca', 'Todos')])
    sheet.append(['Estado', filtros.get('status', 'Todos')])
    sheet.append([])
    sheet.append(['Métrica', 'Valor'])
    sheet.append(['Total de productos', estadisticas['total_productos']])
    sheet.append(['Stock total', float(estadisticas['stock_total'])])
    sheet.append(['Valor total', float(estadisticas['valor_total'])])
    sheet.append(['Productos con stock bajo', estadisticas['stock_bajo']])
    sheet.append([])

    categoria_sheet = workbook.create_sheet('Por categoría')
    categoria_sheet.append(['Categoría', 'Productos', 'Stock'])
    for item in estadisticas['por_categoria']:
        categoria_sheet.append([item['nombre_categoria'], item['total_productos'], float(item['stock_total'])])

    marca_sheet = workbook.create_sheet('Por marca')
    marca_sheet.append(['Marca', 'Productos', 'Stock'])
    for item in estadisticas['por_marca']:
        marca_sheet.append([item['nombre_marca'], item['total_productos'], float(item['stock_total'])])

    stock_sheet = workbook.create_sheet('Stock bajo')
    stock_sheet.append(['Código', 'Producto', 'Stock', 'Precio'])
    for item in estadisticas['stock_bajo_items']:
        stock_sheet.append([item['cod_producto'], item['nombre_producto'], item['stock'], float(item['precio_producto'])])

    for sheet_to_style in [sheet, categoria_sheet, marca_sheet, stock_sheet]:
        if sheet_to_style.title == 'Estadísticas':
            header_row = 8
        else:
            header_row = 1

        for col in range(1, 4):
            cell = sheet_to_style.cell(row=header_row, column=col)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1B365D')
            cell.alignment = Alignment(horizontal='center')

        sheet_to_style.column_dimensions['A'].width = 24
        sheet_to_style.column_dimensions['B'].width = 18
        sheet_to_style.column_dimensions['C'].width = 14
        sheet_to_style.column_dimensions['D'].width = 16

    archivo = os.path.join(tempfile.gettempdir(), f'reporte_productos_estadistico_{int(datetime.now().timestamp())}.xlsx')
    workbook.save(archivo)
    return archivo


def handle_request(ruta):
    modelo = ProductoModelo()

    cod_categoria = int(request.form.get('cod_categoria', 0) or request.args.get('cod_categoria', 0) or 0)
    cod_marca = int(request.form.get('cod_marca', 0) or request.args.get('cod_marca', 0) or 0)
    status = request.form.get('status', 'all') or request.args.get('status', 'all') or 'all'

    filtros = {
        'cod_categoria': cod_categoria,
        'cod_marca': cod_marca,
        'status': status,
    }

    estadisticas = _consultar_estadisticas(modelo, cod_categoria=cod_categoria, cod_marca=cod_marca, status=status)
    categorias, marcas = _listar_catalogos(modelo)

    if request.method == 'POST' and request.form.get('pdf'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_pdf(empresa_data, estadisticas, filtros)
        return send_file(archivo, mimetype='application/pdf', as_attachment=False)

    if request.method == 'POST' and request.form.get('excel'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_excel(empresa_data, estadisticas, filtros)
        return send_file(archivo, as_attachment=True, download_name='reporte_productos_estadistico.xlsx')

    return {
        'estadisticas': estadisticas,
        'categorias': categorias,
        'marcas': marcas,
        'filtros': filtros,
    }
