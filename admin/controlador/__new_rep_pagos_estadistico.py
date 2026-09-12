from flask import request, send_file
from admin.modelo.pagos import PagosModelo
from admin.modelo.empresa import EmpresaModelo
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from datetime import datetime
import tempfile
import os


def _build_filters(desde, hasta, cod_tipo_pago, status):
    condiciones = []
    parametros = []

    if desde:
        condiciones.append('p.fecha_pago >= %s')
        parametros.append(desde)

    if hasta:
        condiciones.append('p.fecha_pago <= %s')
        parametros.append(hasta)

    if cod_tipo_pago and cod_tipo_pago > 0:
        condiciones.append('dp.cod_tipo_pago = %s')
        parametros.append(cod_tipo_pago)

    if status == 'active':
        condiciones.append('p.status = 1')
    elif status == 'anulado':
        condiciones.append('p.status = 0')

    where = ' AND '.join(condiciones) if condiciones else '1=1'
    return where, parametros


def _fetch_scalar(modelo, sql, parametros):
    cursor = modelo.conexion.cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    fila = cursor.fetchone()
    cursor.close()
    return fila


def _fetch_rows(modelo, sql, parametros):
    cursor = modelo.conexion.cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    filas = cursor.fetchall()
    cursor.close()
    return filas


def _consultar_estadisticas(modelo, desde=None, hasta=None, cod_tipo_pago=0, status='all'):
    where, parametros = _build_filters(desde, hasta, cod_tipo_pago, status)

    total_pagos = _fetch_scalar(modelo, f'''
        SELECT COUNT(DISTINCT p.cod_pago) AS total_pagos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
    ''', parametros) or {}

    ingresos = _fetch_scalar(modelo, f'''
        SELECT
            COALESCE(SUM(p.monto), 0) AS ingresos_totales,
            COALESCE(AVG(p.monto), 0) AS pago_promedio
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
    ''', parametros) or {}

    estado = _fetch_scalar(modelo, f'''
        SELECT
            SUM(p.status = 1) AS pagos_activos,
            SUM(p.status = 0) AS pagos_anulados
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
    ''', parametros) or {}

    pagos_por_tipo = _fetch_rows(modelo, f'''
        SELECT
            COALESCE(tp.cod_tipo_pago, 0) AS cod_tipo_pago,
            COALESCE(tp.nombre_tipo_pago, 'Sin tipo') AS nombre_tipo_pago,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN tipo_pago tp ON tp.cod_tipo_pago = dp.cod_tipo_pago
        WHERE {where}
        GROUP BY tp.cod_tipo_pago, tp.nombre_tipo_pago
        ORDER BY ingresos DESC
    ''', parametros)

    pagos_por_moneda = _fetch_rows(modelo, f'''
        SELECT
            COALESCE(m.cod_moneda, 0) AS cod_moneda,
            COALESCE(m.nombre_moneda, 'Bs') AS nombre_moneda,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN moneda m ON m.cod_moneda = COALESCE(dp.cod_moneda, 1)
        WHERE {where}
        GROUP BY m.cod_moneda, m.nombre_moneda
        ORDER BY ingresos DESC
    ''', parametros)

    pagos_por_mes = _fetch_rows(modelo, f'''
        SELECT
            DATE_FORMAT(p.fecha_pago, '%%b %%Y') AS periodo,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
        GROUP BY YEAR(p.fecha_pago), MONTH(p.fecha_pago)
        ORDER BY YEAR(p.fecha_pago), MONTH(p.fecha_pago)
        LIMIT 12
    ''', parametros)

    top_clientes = _fetch_rows(modelo, f'''
        SELECT
            u.cedula,
            u.nombre_usuario AS nombre_cliente,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        INNER JOIN servicio s ON p.cod_servicio = s.cod_servicio
        INNER JOIN seguridad.usuario u ON s.cedula = u.cedula
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
        GROUP BY u.cedula, u.nombre_usuario
        ORDER BY ingresos DESC
        LIMIT 5
    ''', parametros)

    ultimos_pagos = _fetch_rows(modelo, f'''
        SELECT
            p.cod_pago,
            p.fecha_pago,
            p.monto,
            p.status,
            u.nombre_usuario AS cliente,
            COALESCE(tp.nombre_tipo_pago, 'N/A') AS tipo_pago,
            COALESCE(m.nombre_moneda, 'Bs') AS moneda
        FROM pago p
        INNER JOIN servicio s ON p.cod_servicio = s.cod_servicio
        INNER JOIN seguridad.usuario u ON s.cedula = u.cedula
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN tipo_pago tp ON tp.cod_tipo_pago = dp.cod_tipo_pago
        LEFT JOIN moneda m ON m.cod_moneda = COALESCE(dp.cod_moneda, 1)
        WHERE {where}
        ORDER BY p.fecha_pago DESC
        LIMIT 10
    ''', parametros)

    return {
        'total_pagos': int(total_pagos.get('total_pagos', 0)),
        'ingresos_totales': float(ingresos.get('ingresos_totales', 0.0)),
        'pago_promedio': float(ingresos.get('pago_promedio', 0.0)),
        'pagos_activos': int(estado.get('pagos_activos', 0)),
        'pagos_anulados': int(estado.get('pagos_anulados', 0)),
        'pagos_por_tipo': pagos_por_tipo,
        'pagos_por_moneda': pagos_por_moneda,
        'pagos_por_mes': pagos_por_mes,
        'top_clientes': top_clientes,
        'ultimos_pagos': ultimos_pagos,
    }


def _generar_pdf(empresa, estadisticas, filtros):
    archivo = os.path.join(tempfile.gettempdir(), f'reporte_pagos_estadistico_{int(datetime.now().timestamp())}.pdf')
    doc = SimpleDocTemplate(archivo, pagesize=LETTER, rightMargin=18, leftMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph('Reporte Estadístico de Pagos', styles['Title']))
    elements.append(Spacer(1, 8))
    elements.append(Paragraph(f"Emitido: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles['Normal']))
    elements.append(Spacer(1, 12))

    filtros_data = [
        ['Desde', filtros.get('desde', '')],
        ['Hasta', filtros.get('hasta', '')],
        ['Tipo de Pago', filtros.get('cod_tipo_pago', 'Todos')],
        ['Estado', filtros.get('status', 'Todos')],
    ]
    filtros_table = Table(filtros_data, hAlign='LEFT', colWidths=[120, 260])
    filtros_table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
    ]))
    elements.extend([filtros_table, Spacer(1, 12)])

    metrics_data = [
        ['Métrica', 'Valor'],
        ['Total de pagos', str(estadisticas['total_pagos'])],
        ['Ingresos totales', f"{estadisticas['ingresos_totales']:.2f}"],
        ['Pago promedio', f"{estadisticas['pago_promedio']:.2f}"],
        ['Pagos anulados', str(estadisticas['pagos_anulados'])],
    ]
    metrics_table = Table(metrics_data, hAlign='LEFT', colWidths=[180, 160])
    metrics_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4a69bd')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
    ]))
    elements.extend([metrics_table, Spacer(1, 12)])

    if estadisticas['pagos_por_tipo']:
        elements.append(Paragraph('Pagos por tipo', styles['Heading3']))
        tipo_data = [['Tipo', 'Pagos', 'Ingresos']]
        for item in estadisticas['pagos_por_tipo']:
            tipo_data.append([item['nombre_tipo_pago'], item['total_pagos'], f"{item['ingresos']:.2f}"])
        tipo_table = Table(tipo_data, hAlign='LEFT', colWidths=[200, 80, 120])
        tipo_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ]))
        elements.extend([tipo_table, Spacer(1, 12)])

    if estadisticas['pagos_por_moneda']:
        elements.append(Paragraph('Pagos por moneda', styles['Heading3']))
        moneda_data = [['Moneda', 'Pagos', 'Ingresos']]
        for item in estadisticas['pagos_por_moneda']:
            moneda_data.append([item['nombre_moneda'], item['total_pagos'], f"{item['ingresos']:.2f}"])
        moneda_table = Table(moneda_data, hAlign='LEFT', colWidths=[200, 80, 120])
        moneda_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ]))
        elements.extend([moneda_table, Spacer(1, 12)])

    if estadisticas['top_clientes']:
        elements.append(Paragraph('Top clientes', styles['Heading3']))
        clientes_data = [['Cliente', 'Pagos', 'Ingresos']]
        for cliente in estadisticas['top_clientes']:
            clientes_data.append([cliente['nombre_cliente'], cliente['total_pagos'], f"{cliente['ingresos']:.2f}"])
        clientes_table = Table(clientes_data, hAlign='LEFT', colWidths=[220, 70, 110])
        clientes_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ]))
        elements.extend([clientes_table, Spacer(1, 12)])

    doc.build(elements)
    return archivo


def _generar_excel(empresa, estadisticas, filtros):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Estadísticas'

    sheet.append(['Reporte Estadístico de Pagos'])
    sheet.append(['Emitido', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    sheet.append([])
    sheet.append(['Desde', filtros.get('desde', '')])
    sheet.append(['Hasta', filtros.get('hasta', '')])
    sheet.append(['Tipo de Pago', filtros.get('cod_tipo_pago', 'Todos')])
    sheet.append(['Estado', filtros.get('status', 'Todos')])
    sheet.append([])
    sheet.append(['Métrica', 'Valor'])
    sheet.append(['Total de pagos', estadisticas['total_pagos']])
    sheet.append(['Ingresos totales', float(estadisticas['ingresos_totales'])])
    sheet.append(['Pago promedio', float(estadisticas['pago_promedio'])])
    sheet.append(['Pagos anulados', estadisticas['pagos_anulados']])
    sheet.append([])

    tipo_sheet = workbook.create_sheet('Pagos por tipo')
    tipo_sheet.append(['Tipo', 'Pagos', 'Ingresos'])
    for item in estadisticas['pagos_por_tipo']:
        tipo_sheet.append([item['nombre_tipo_pago'], item['total_pagos'], float(item['ingresos'])])

    moneda_sheet = workbook.create_sheet('Pagos por moneda')
    moneda_sheet.append(['Moneda', 'Pagos', 'Ingresos'])
    for item in estadisticas['pagos_por_moneda']:
        moneda_sheet.append([item['nombre_moneda'], item['total_pagos'], float(item['ingresos'])])

    top_sheet = workbook.create_sheet('Top clientes')
    top_sheet.append(['Cliente', 'Pagos', 'Ingresos'])
    for item in estadisticas['top_clientes']:
        top_sheet.append([item['nombre_cliente'], item['total_pagos'], float(item['ingresos'])])

    archivo = os.path.join(tempfile.gettempdir(), f'reporte_pagos_estadistico_{int(datetime.now().timestamp())}.xlsx')
    workbook.save(archivo)
    return archivo


def handle_request(ruta):
    modelo = PagosModelo()

    desde = request.form.get('desde') or request.args.get('desde')
    hasta = request.form.get('hasta') or request.args.get('hasta')
    cod_tipo_pago = int(request.form.get('cod_tipo_pago', 0) or 0)
    status = request.form.get('status', 'all')

    filtros = {
        'desde': desde or '',
        'hasta': hasta or '',
        'cod_tipo_pago': cod_tipo_pago,
        'status': status,
    }

    estadisticas = _consultar_estadisticas(
        modelo,
        desde=desde,
        hasta=hasta,
        cod_tipo_pago=cod_tipo_pago,
        status=status,
    )
    tipopagos = modelo.listar_tipos_pago()

    if request.method == 'POST' and request.form.get('pdf'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_pdf(empresa_data, estadisticas, filtros)
        return send_file(archivo, as_attachment=True, download_name='reporte_pagos_estadistico.pdf')

    if request.method == 'POST' and request.form.get('excel'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_excel(empresa_data, estadisticas, filtros)
        return send_file(archivo, as_attachment=True, download_name='reporte_pagos_estadistico.xlsx')

    return {
        'estadisticas': estadisticas,
        'tipopagos': tipopagos,
        'filtros': filtros,
    }
