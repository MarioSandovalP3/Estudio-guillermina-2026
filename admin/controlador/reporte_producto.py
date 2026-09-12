from flask import request, send_file
from admin.modelo.producto import ProductoModelo
from admin.modelo.empresa import EmpresaModelo
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
import tempfile
import os


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 9)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 12 * mm, 'Sistema de Gestión | Reporte generado automáticamente')
    canvas.drawRightString(195 * mm, 12 * mm, f'Página {doc.page}')
    canvas.restoreState()


def borde(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.lightgrey)
    canvas.setLineWidth(1)
    canvas.rect(20 * mm, 20 * mm, LETTER[0] - 40 * mm, LETTER[1] - 40 * mm)
    canvas.restoreState()


def _build_filters(cod_categoria, cod_marca, status, vencimiento='all', stock='all'):
    condiciones = []
    parametros = []

    if cod_categoria and cod_categoria > 0:
        condiciones.append('p.cod_categoria = %s')
        parametros.append(cod_categoria)

    if cod_marca and cod_marca > 0:
        condiciones.append('p.cod_marca = %s')
        parametros.append(cod_marca)

    if status == 'active':
        condiciones.append('p.status = 1')
    elif status == 'inactive':
        condiciones.append('p.status = 0')

    if vencimiento == 'expired':
        condiciones.append('dp.fecha_vencimiento < CURDATE()')
    elif vencimiento == 'next_30':
        condiciones.append('dp.fecha_vencimiento BETWEEN CURDATE() AND DATE_ADD(CURDATE(), INTERVAL 30 DAY)')
    elif vencimiento == 'after_30':
        condiciones.append('dp.fecha_vencimiento > DATE_ADD(CURDATE(), INTERVAL 30 DAY)')
    elif vencimiento == 'without_date':
        condiciones.append('dp.fecha_vencimiento IS NULL')

    if stock == 'empty':
        condiciones.append('p.stock = 0')
    elif stock == 'low':
        condiciones.append('p.stock BETWEEN 1 AND 5')
    elif stock == 'available':
        condiciones.append('p.stock > 5')

    where = ' AND '.join(condiciones) if condiciones else '1=1'
    return where, parametros


def _fetch_rows(modelo, sql, parametros):
    cursor = modelo.conn().cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    filas = cursor.fetchall()
    cursor.close()
    return filas


def _consultar_productos(modelo, cod_categoria=0, cod_marca=0, status='all', vencimiento='all', stock='all'):
    where, parametros = _build_filters(cod_categoria, cod_marca, status, vencimiento, stock)
    sql = f'''
        SELECT
            p.cod_producto,
            p.nombre_producto,
            p.precio_producto,
            p.stock,
            p.status,
            c.nombre_categoria,
            m.nombre_marca,
            u.nombre_unidad,
            COALESCE(pr.presentacion, 'N/A') AS nombre_presentacion,
            COALESCE(me.medida, 'N/A') AS nombre_medida,
            COALESCE(dp.nombre_tipo_producto, 'Sin tipo') AS nombre_tipo_producto,
            dp.fecha_vencimiento,
            DATEDIFF(dp.fecha_vencimiento, CURDATE()) AS dias_para_vencer,
            (p.precio_producto * p.stock) AS valor_inventario,
            CASE
                WHEN p.stock = 0 THEN 'Agotado'
                WHEN p.stock <= 5 THEN 'Stock bajo'
                ELSE 'Disponible'
            END AS nivel_stock
        FROM producto p
        INNER JOIN categoria c ON p.cod_categoria = c.cod_categoria
        INNER JOIN marca m ON p.cod_marca = m.cod_marca
        INNER JOIN unidad u ON p.cod_unidad = u.cod_unidad
        LEFT JOIN presentacion pr ON u.cod_presentacion = pr.cod_presentacion
        LEFT JOIN medida me ON u.cod_medida = me.cod_medida
        LEFT JOIN (
            SELECT
                dp.cod_producto,
                MIN(dp.fecha_vencimiento) AS fecha_vencimiento,
                GROUP_CONCAT(DISTINCT tp.nombre_tipo_producto ORDER BY tp.nombre_tipo_producto SEPARATOR ', ') AS nombre_tipo_producto
            FROM detalle_producto dp
            LEFT JOIN tipo_producto tp ON dp.cod_tipo_producto = tp.cod_tipo_producto
            GROUP BY dp.cod_producto
        ) dp ON p.cod_producto = dp.cod_producto
        WHERE {where}
        ORDER BY p.cod_producto DESC
    '''
    return _fetch_rows(modelo, sql, parametros)


def _consultar_resumen(modelo, cod_categoria=0, cod_marca=0, status='all', vencimiento='all', stock='all'):
    where, parametros = _build_filters(cod_categoria, cod_marca, status, vencimiento, stock)
    resumen = _fetch_rows(modelo, f'''
        SELECT
            COUNT(*) AS total_productos,
            COALESCE(SUM(p.stock), 0) AS stock_total,
            COALESCE(SUM(p.precio_producto * p.stock), 0) AS valor_inventario,
            COALESCE(AVG(p.precio_producto), 0) AS precio_promedio,
            SUM(p.stock = 0) AS agotados,
            SUM(p.stock BETWEEN 1 AND 5) AS stock_bajo,
            SUM(p.status = 1) AS activos,
            SUM(dp.fecha_vencimiento BETWEEN CURDATE() AND DATE_ADD(CURDATE(), INTERVAL 30 DAY)) AS vencen_30_dias
        FROM producto p
        LEFT JOIN (
            SELECT cod_producto, MIN(fecha_vencimiento) AS fecha_vencimiento
            FROM detalle_producto
            GROUP BY cod_producto
        ) dp ON p.cod_producto = dp.cod_producto
        WHERE {where}
    ''', parametros)
    return resumen[0] if resumen else {}


def _listar_catalogos(modelo):
    cursor = modelo.conn().cursor(dictionary=True)
    cursor.execute('SELECT cod_categoria, nombre_categoria FROM categoria ORDER BY nombre_categoria')
    categorias = cursor.fetchall()
    cursor.execute('SELECT cod_marca, nombre_marca FROM marca ORDER BY nombre_marca')
    marcas = cursor.fetchall()
    cursor.close()
    return categorias, marcas


def _generar_pdf(empresa, listado, filtros, resumen):
    archivo = os.path.join(tempfile.gettempdir(), f'reporte_productos_{int(datetime.now().timestamp())}.pdf')
    doc = SimpleDocTemplate(
        archivo,
        pagesize=LETTER,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm
    )
    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'titulo',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1B365D'),
        alignment=1,
        spaceAfter=6
    )

    normal_style = ParagraphStyle(
        'normal',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14
    )

    elementos = []

    if empresa.get('logo'):
        base_dir = os.path.dirname(os.path.abspath(__file__))
        logo_path = os.path.abspath(os.path.join(base_dir, '..', '..', 'static', 'assets', 'img', empresa['logo']))
        if os.path.exists(logo_path):
            imagen = Image(logo_path, width=50 * mm, height=30 * mm)
            imagen.hAlign = 'LEFT'
            elementos.append(imagen)
            elementos.append(Spacer(1, 8))

    elementos.append(Paragraph('Estudio MT - Salón & Estética', titulo_style))
    elementos.append(Paragraph(f"Dirección: {empresa.get('direccion','')}" if empresa.get('direccion') else 'Dirección: N/A', normal_style))
    elementos.append(Paragraph(f"Teléfono: {empresa.get('telefono','')}" if empresa.get('telefono') else 'Teléfono: N/A', normal_style))
    elementos.append(Paragraph(f"Email: {empresa.get('email','')}" if empresa.get('email') else 'Email: N/A', normal_style))
    elementos.append(Spacer(1, 12))

    elementos.append(Paragraph('<b>REPORTE DE PRODUCTOS</b>', styles['Heading2']))
    elementos.append(Paragraph(f"Fecha: {datetime.now().strftime('%d-%m-%Y')}", normal_style))
    elementos.append(Spacer(1, 12))

    if filtros:
        filtros_data = [
            ['Categoría', filtros.get('cod_categoria', 'Todos')],
            ['Marca', filtros.get('cod_marca', 'Todos')],
            ['Estado', filtros.get('status', 'Todos')],
            ['Vencimiento', filtros.get('vencimiento', 'Todos')],
            ['Cantidad de stock', filtros.get('stock', 'Todos')],
        ]
        filtros_table = Table(filtros_data, hAlign='LEFT', colWidths=[30 * mm, 105 * mm])
        filtros_table.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 3),
            ('RIGHTPADDING', (0, 0), (-1, -1), 3),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ]))
        elementos.extend([filtros_table, Spacer(1, 12)])

    resumen_data = [
        ['Indicador', 'Resultado'],
        ['Productos filtrados', str(resumen.get('total_productos', 0))],
        ['Unidades en inventario', str(resumen.get('stock_total', 0))],
        ['Valor del inventario', f"{float(resumen.get('valor_inventario', 0) or 0):.2f}"],
        ['Agotados / stock bajo', f"{resumen.get('agotados', 0) or 0} / {resumen.get('stock_bajo', 0) or 0}"],
        ['Vencen en 30 días', str(resumen.get('vencen_30_dias', 0) or 0)],
    ]
    resumen_table = Table(resumen_data, hAlign='LEFT', colWidths=[80 * mm, 65 * mm])
    resumen_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey]),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
    ]))
    elementos.extend([resumen_table, Spacer(1, 12)])

    data = [['Código', 'Producto', 'Categoría', 'Marca', 'Tipo', 'Stock', 'Precio', 'Valor', 'Nivel', 'Vence', 'Estado']]
    for producto in listado:
        estado = 'Activo' if producto['status'] == 1 else 'Inactivo'
        data.append([
            producto['cod_producto'],
            producto['nombre_producto'],
            producto['nombre_categoria'] or 'N/A',
            producto['nombre_marca'] or 'N/A',
            producto.get('nombre_tipo_producto') or 'Sin tipo',
            str(producto['stock']),
            f"{producto['precio_producto']:.2f}",
            f"{float(producto.get('valor_inventario', 0) or 0):.2f}",
            producto.get('nivel_stock') or 'N/A',
            str(producto.get('fecha_vencimiento') or 'N/A'),
            estado,
        ])

    tabla = Table(data, repeatRows=1, hAlign='LEFT', colWidths=[9 * mm, 22 * mm, 16 * mm, 14 * mm, 18 * mm, 9 * mm, 12 * mm, 14 * mm, 16 * mm, 16 * mm, 12 * mm])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (0, 1), (-1, -1), 'LEFT'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 2),
        ('RIGHTPADDING', (0, 0), (-1, -1), 2),
        ('TOPPADDING', (0, 0), (-1, -1), 1),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 1),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
    ]))
    elementos.append(tabla)

    doc.build(
        elementos,
        onFirstPage=lambda c, d: (borde(c, d), footer(c, d)),
        onLaterPages=lambda c, d: (borde(c, d), footer(c, d))
    )
    return archivo


def _generar_excel(empresa, listado, filtros, resumen):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Productos'

    sheet.merge_cells('A1:G1')
    sheet['A1'] = empresa.get('nombre', 'Estudio MT')
    sheet['A1'].font = Font(bold=True, size=14)
    sheet['A1'].alignment = Alignment(horizontal='center')

    sheet['A2'] = empresa.get('direccion', '')
    sheet['A3'] = empresa.get('telefono', '')
    sheet['A4'] = empresa.get('email', '')
    sheet['A5'] = datetime.now().strftime('%d-%m-%Y')

    sheet.append([])
    sheet.append(['Reporte de Productos'])
    sheet.append(['Emitido', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    sheet.append([])
    sheet.append(['Categoría', filtros.get('cod_categoria', 'Todos')])
    sheet.append(['Marca', filtros.get('cod_marca', 'Todos')])
    sheet.append(['Estado', filtros.get('status', 'Todos')])
    sheet.append(['Vencimiento', filtros.get('vencimiento', 'Todos')])
    sheet.append(['Cantidad de stock', filtros.get('stock', 'Todos')])
    sheet.append([])
    sheet.append(['Indicador', 'Resultado'])
    sheet.append(['Productos filtrados', resumen.get('total_productos', 0)])
    sheet.append(['Unidades en inventario', resumen.get('stock_total', 0)])
    sheet.append(['Valor del inventario', float(resumen.get('valor_inventario', 0) or 0)])
    sheet.append(['Agotados', resumen.get('agotados', 0) or 0])
    sheet.append(['Stock bajo', resumen.get('stock_bajo', 0) or 0])
    sheet.append(['Vencen en 30 días', resumen.get('vencen_30_dias', 0) or 0])
    sheet.append([])
    detail_header_row = sheet.max_row + 1
    sheet.append(['Código', 'Producto', 'Categoría', 'Marca', 'Unidad', 'Presentación', 'Medida', 'Tipo', 'Stock', 'Precio', 'Valor inventario', 'Nivel', 'Vencimiento', 'Días para vencer', 'Estado'])

    for producto in listado:
        sheet.append([
            producto['cod_producto'],
            producto['nombre_producto'],
            producto['nombre_categoria'] or 'N/A',
            producto['nombre_marca'] or 'N/A',
            producto.get('nombre_unidad') or 'N/A',
            producto.get('nombre_presentacion') or 'N/A',
            producto.get('nombre_medida') or 'N/A',
            producto.get('nombre_tipo_producto') or 'Sin tipo',
            producto['stock'],
            float(producto['precio_producto']),
            float(producto.get('valor_inventario', 0) or 0),
            producto.get('nivel_stock') or 'N/A',
            producto.get('fecha_vencimiento') or 'N/A',
            producto.get('dias_para_vencer') if producto.get('dias_para_vencer') is not None else 'N/A',
            'Activo' if producto['status'] == 1 else 'Inactivo',
        ])

    for col in range(1, 16):
        cell = sheet.cell(row=detail_header_row, column=col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1B365D')
        cell.alignment = Alignment(horizontal='center')

    sheet.column_dimensions['A'].width = 10
    sheet.column_dimensions['B'].width = 24
    sheet.column_dimensions['C'].width = 18
    sheet.column_dimensions['D'].width = 16
    sheet.column_dimensions['E'].width = 14
    sheet.column_dimensions['F'].width = 16
    sheet.column_dimensions['G'].width = 14
    sheet.column_dimensions['H'].width = 14
    sheet.column_dimensions['I'].width = 10
    sheet.column_dimensions['J'].width = 12
    sheet.column_dimensions['K'].width = 16
    sheet.column_dimensions['L'].width = 16
    sheet.column_dimensions['M'].width = 16
    sheet.column_dimensions['N'].width = 16
    sheet.column_dimensions['O'].width = 12

    archivo = os.path.join(tempfile.gettempdir(), f'reporte_productos_{int(datetime.now().timestamp())}.xlsx')
    workbook.save(archivo)
    return archivo


def handle_request(ruta):
    modelo = ProductoModelo()

    cod_categoria_raw = request.values.get('cod_categoria', 0)
    cod_marca_raw = request.values.get('cod_marca', 0)
    status = request.values.get('status', 'all') or 'all'
    vencimiento = request.values.get('vencimiento', 'all') or 'all'
    stock = request.values.get('stock', 'all') or 'all'

    try:
        cod_categoria = int(cod_categoria_raw or 0)
    except (TypeError, ValueError):
        cod_categoria = 0

    try:
        cod_marca = int(cod_marca_raw or 0)
    except (TypeError, ValueError):
        cod_marca = 0

    filtros = {
        'cod_categoria': cod_categoria,
        'cod_marca': cod_marca,
        'status': status,
        'vencimiento': vencimiento,
        'stock': stock,
    }

    registro = _consultar_productos(modelo, cod_categoria=cod_categoria, cod_marca=cod_marca, status=status,
                                    vencimiento=vencimiento, stock=stock)
    resumen = _consultar_resumen(modelo, cod_categoria=cod_categoria, cod_marca=cod_marca, status=status,
                                 vencimiento=vencimiento, stock=stock)
    categorias, marcas = _listar_catalogos(modelo)

    if request.method == 'POST' and request.form.get('pdf'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_pdf(empresa_data, registro, filtros, resumen)
        return send_file(archivo, mimetype='application/pdf', as_attachment=False)

    if request.method == 'POST' and request.form.get('excel'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_excel(empresa_data, registro, filtros, resumen)
        return send_file(archivo, as_attachment=True, download_name='reporte_productos.xlsx')

    return {
        'registro': registro,
        'resumen': resumen,
        'categorias': categorias,
        'marcas': marcas,
        'filtros': filtros,
    }
