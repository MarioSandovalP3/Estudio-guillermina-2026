from flask import request, send_file

from admin.modelo.reservas import ReservasModelo
from admin.modelo.empresa import EmpresaModelo
from admin.modelo.panel import PanelModelo

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Table,
    TableStyle,
    Spacer,
    Image
)
from reportlab.graphics.shapes import Drawing, Line
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

from datetime import datetime
import tempfile
import os


def footer(canvas, doc):

    canvas.saveState()

    canvas.setFont(
        'Helvetica',
        9
    )

    canvas.setFillColor(
        colors.HexColor("#777777")
    )

    canvas.drawString(
        15 * mm,
        12 * mm,
        'Sistema de Gestión | Reporte generado automáticamente'
    )

    canvas.drawRightString(
        195 * mm,
        12 * mm,
        f'Página {doc.page}'
    )

    canvas.restoreState()


def borde(canvas, doc):

    canvas.saveState()

    canvas.setStrokeColor(
        colors.HexColor("#D4AF37")
    )

    canvas.setLineWidth(
        1.2
    )

    canvas.rect(
        10 * mm,
        10 * mm,
        LETTER[0] - 20 * mm,
        LETTER[1] - 20 * mm
    )

    canvas.restoreState()


def _build_filters(
    desde,
    hasta,
    status
):

    condiciones = []
    parametros = []

    if desde:

        condiciones.append(
            's.fecha >= %s'
        )

        parametros.append(
            desde
        )

    if hasta:

        condiciones.append(
            's.fecha <= %s'
        )

        parametros.append(
            hasta
        )

    if status == 'active':

        condiciones.append(
            's.status = 1'
        )

    elif status == 'Cancelada':

        condiciones.append(
            's.status = 0'
        )

    where = (
        ' AND '.join(condiciones)
        if condiciones
        else '1=1'
    )

    return where, parametros


def _fetch_rows(
    modelo,
    sql,
    parametros
):

    cursor = modelo.conexion.cursor(
        dictionary=True
    )

    cursor.execute(
        sql,
        tuple(parametros)
    )

    filas = cursor.fetchall()

    cursor.close()

    return filas


def _fetch_scalar(
    modelo,
    sql,
    parametros
):

    cursor = modelo.conexion.cursor(
        dictionary=True
    )

    cursor.execute(
        sql,
        tuple(parametros)
    )

    fila = cursor.fetchone()

    cursor.close()

    return fila


def _consultar_estadisticas(
    modelo,
    desde=None,
    hasta=None,
    status='all'
):

    where, parametros = _build_filters(
        desde,
        hasta,
        status
    )

    total = _fetch_scalar(
        modelo,
        f'''
        SELECT
            COUNT(*) AS total_reservas,
            SUM(s.status = 1) AS reservas_activas,
            SUM(s.status = 0) AS reservas_canceladas
        FROM sac.servicio s
        WHERE {where}
        ''',
        parametros
    ) or {}

    reservas_por_servicio = _fetch_rows(
        modelo,
        f'''
        SELECT
            COALESCE(
                ts.nombre_servicio,
                'Sin servicio'
            ) AS servicio,

            COUNT(DISTINCT s.cod_servicio)
                AS total_reservas

        FROM sac.servicio s

        LEFT JOIN sac.detalle_servicio ds
            ON ds.cod_servicio =
               s.cod_servicio

        LEFT JOIN sac.tipo_servicio ts
            ON ts.cod_tipo_servicio =
               ds.cod_tipo_servicio

        WHERE {where}

        GROUP BY
            ts.nombre_servicio

        ORDER BY
            total_reservas DESC
        ''',
        parametros
    )

    reservas_por_especialista = _fetch_rows(
        modelo,
        f'''
        SELECT
            COALESCE(
                e.especialista,
                CONCAT(
                    u2.nombre_usuario,
                    ' ',
                    u2.apellido_usuario
                )
            ) AS especialista,

            COUNT(DISTINCT s.cod_servicio)
                AS total_reservas

        FROM sac.servicio s

        LEFT JOIN sac.especialista e
            ON s.cedula_especialista =
               e.cedula_especialista

        LEFT JOIN seguridad.usuario u2
            ON u2.cedula =
               e.cedula_especialista

        WHERE {where}

        GROUP BY
            especialista

        ORDER BY
            total_reservas DESC
        ''',
        parametros
    )

    top_clientes = _fetch_rows(
        modelo,
        f'''
        SELECT
            u.cedula,

            CONCAT(
                u.nombre_usuario,
                ' ',
                u.apellido_usuario
            ) AS cliente,

            COUNT(DISTINCT s.cod_servicio)
                AS total_reservas

        FROM sac.servicio s

        INNER JOIN seguridad.usuario u
            ON s.cedula = u.cedula

        WHERE {where}

        GROUP BY
            u.cedula,
            u.nombre_usuario,
            u.apellido_usuario

        ORDER BY
            total_reservas DESC

        LIMIT 5
        ''',
        parametros
    )

    reservas_por_mes = _fetch_rows(
        modelo,
        f'''
        SELECT
            DATE_FORMAT(
                s.fecha,
                '%%b %%Y'
            ) AS periodo,

            COUNT(*) AS total_reservas

        FROM sac.servicio s

        WHERE {where}

        GROUP BY
            YEAR(s.fecha),
            MONTH(s.fecha)

        ORDER BY
            YEAR(s.fecha),
            MONTH(s.fecha)

        LIMIT 12
        ''',
        parametros
    )

    ultimas_reservas = _fetch_rows(
        modelo,
        f'''
        SELECT
            s.cod_servicio,
            s.fecha AS fecha_reserva,
            s.hora AS hora_reserva,
            s.status,

            CONCAT(
                u.nombre_usuario,
                ' ',
                u.apellido_usuario
            ) AS cliente,

            COALESCE(
                e.especialista,
                'N/A'
            ) AS especialista

        FROM sac.servicio s

        INNER JOIN seguridad.usuario u
            ON s.cedula = u.cedula

        LEFT JOIN sac.especialista e
            ON s.cedula_especialista =
               e.cedula_especialista

        WHERE {where}

        ORDER BY
            s.fecha DESC,
            s.hora DESC

        LIMIT 10
        ''',
        parametros
    )

    return {
        'total_reservas': int(
            total.get(
                'total_reservas',
                0
            ) or 0
        ),

        'reservas_activas': int(
            total.get(
                'reservas_activas',
                0
            ) or 0
        ),

        'reservas_inactivas': int(
            total.get(
                'reservas_inactivas',
                0
            ) or 0
        ),

        'reservas_por_servicio':
            reservas_por_servicio,

        'reservas_por_especialista':
            reservas_por_especialista,

        'top_clientes':
            top_clientes,

        'reservas_por_mes':
            reservas_por_mes,

        'ultimas_reservas':
            ultimas_reservas
    }


# ==============================================================
# ESTADÍSTICAS INDIVIDUALES DEL CLIENTE
# ==============================================================
#
# STATUS:
#
# 0 = CANCELADA
# 1 = ACTIVA
# 2 = NO SE MUESTRA
#
# ==============================================================

def _consultar_estadisticas_cliente(
    modelo,
    cedula_cliente
):

    if not cedula_cliente:

        return {
            "cliente": {},
            "total_reservas": 0,
            "especialista_mas_atendio": "—",
            "servicio_mas_solicitado": "—",
            "reservas": []
        }

    cursor = modelo.conexion.cursor(
        dictionary=True
    )

    # ==========================================================
    # BUSCAR CLIENTE
    # ==========================================================

    cursor.execute(
        """
        SELECT
            u.cedula,
            u.nombre_usuario,
            u.apellido_usuario

        FROM seguridad.usuario u

        INNER JOIN seguridad.rol r
            ON r.cod_rol = u.cod_rol

        WHERE u.cedula = %s
          AND r.cod_rol = 3

        LIMIT 1
        """,
        (cedula_cliente,)
    )

    cliente = cursor.fetchone()

    if not cliente:

        cursor.close()

        return {
            "cliente": {},
            "total_reservas": 0,
            "especialista_mas_atendio": "—",
            "servicio_mas_solicitado": "—",
            "reservas": []
        }


    # ==========================================================

    cursor.execute(
        """
        SELECT
            COUNT(
                DISTINCT s.cod_servicio
            ) AS total_reservas

        FROM sac.servicio s

        WHERE s.cedula = %s
          AND s.status IN (0, 1)
        """,
        (cedula_cliente,)
    )

    resultado = cursor.fetchone()

    total_reservas = (
        resultado["total_reservas"]
        if resultado
        else 0
    )

    # ==========================================================
    # ESPECIALISTA QUE MÁS ATENDIÓ
    #
    # SOLO RESERVAS STATUS 0 Y 1
    # ==========================================================

    cursor.execute(
        """
        SELECT

            CONCAT(
                e.especialista,
                ' — ',
                u.nombre_usuario,
                ' ',
                u.apellido_usuario
            ) AS especialista,

            COUNT(
                DISTINCT s.cod_servicio
            ) AS cantidad

        FROM sac.servicio s

        INNER JOIN sac.especialista e
            ON e.cedula_especialista =
               s.cedula_especialista

        INNER JOIN seguridad.usuario u
            ON u.cedula =
               e.cedula_especialista

        WHERE s.cedula = %s
          AND s.status IN (0, 1)

        GROUP BY
            e.cedula_especialista,
            e.especialista,
            u.nombre_usuario,
            u.apellido_usuario

        ORDER BY
            cantidad DESC,
            u.nombre_usuario ASC,
            u.apellido_usuario ASC

        LIMIT 1
        """,
        (cedula_cliente,)
    )

    resultado = cursor.fetchone()

    especialista_mas_atendio = (
        resultado["especialista"]
        if resultado
        else "—"
    )

    # ==========================================================
    # SERVICIO MÁS SOLICITADO
    #
    # SOLO RESERVAS STATUS 0 Y 1
    # ==========================================================

    cursor.execute(
        """
        SELECT

            TRIM(
                ts.nombre_servicio
            ) AS servicio,

            COUNT(
                DISTINCT s.cod_servicio
            ) AS cantidad

        FROM sac.servicio s

        INNER JOIN sac.detalle_servicio ds
            ON ds.cod_servicio =
               s.cod_servicio

        INNER JOIN sac.tipo_servicio ts
            ON ts.cod_tipo_servicio =
               ds.cod_tipo_servicio

        WHERE s.cedula = %s
          AND s.status IN (0, 1)

        GROUP BY
            ts.cod_tipo_servicio,
            ts.nombre_servicio

        ORDER BY
            cantidad DESC,
            ts.nombre_servicio ASC

        LIMIT 1
        """,
        (cedula_cliente,)
    )

    resultado = cursor.fetchone()

    servicio_mas_solicitado = (
        resultado["servicio"]
        if resultado
        else "—"
    )


    cursor.execute(
        """
        SELECT

            s.fecha AS fecha,

            s.hora AS hora,

            CONCAT(
                e.especialista,
                ' — ',
                u.nombre_usuario,
                ' ',
                u.apellido_usuario
            ) AS especialista,

            GROUP_CONCAT(
                DISTINCT TRIM(
                    ts.nombre_servicio
                )
                ORDER BY
                    TRIM(
                        ts.nombre_servicio
                    )
                SEPARATOR ', '
            ) AS servicio,

            CASE

                WHEN p.cod_promo IS NOT NULL
                    THEN CONCAT(
                        p.descuento,
                        '%'
                    )

                ELSE 'No'

            END AS promocion,

            CASE

                WHEN s.status = 0
                    THEN 'Cancelada'

                WHEN s.status = 1
                    THEN 'Activa'

            END AS estado

        FROM sac.servicio s

        INNER JOIN sac.especialista e
            ON e.cedula_especialista =
               s.cedula_especialista

        INNER JOIN seguridad.usuario u
            ON u.cedula =
               e.cedula_especialista

        LEFT JOIN sac.detalle_servicio ds
            ON ds.cod_servicio =
               s.cod_servicio

        LEFT JOIN sac.tipo_servicio ts
            ON ts.cod_tipo_servicio =
               ds.cod_tipo_servicio

        LEFT JOIN sac.promocion p
            ON p.cod_promo =
               s.cod_promo

        WHERE s.cedula = %s
          AND s.status IN (0, 1)

        GROUP BY

            s.cod_servicio,
            s.fecha,
            s.hora,
            e.especialista,
            u.nombre_usuario,
            u.apellido_usuario,
            p.cod_promo,
            p.descuento,
            s.status

        ORDER BY
            s.fecha DESC,
            s.hora DESC
        """,
        (cedula_cliente,)
    )

    reservas = cursor.fetchall()



    for reserva in reservas:

        fecha = reserva.get(
            "fecha"
        )

        if fecha:

            if hasattr(
                fecha,
                "strftime"
            ):

                reserva["fecha"] = (
                    fecha.strftime(
                        "%d/%m/%Y"
                    )
                )

            else:

                fecha_texto = str(
                    fecha
                )

                if len(fecha_texto) >= 10:

                    try:

                        reserva["fecha"] = (
                            fecha_texto[8:10]
                            + "/"
                            + fecha_texto[5:7]
                            + "/"
                            + fecha_texto[0:4]
                        )

                    except Exception:

                        reserva["fecha"] = (
                            fecha_texto
                        )

                else:

                    reserva["fecha"] = (
                        fecha_texto
                    )

        else:

            reserva["fecha"] = "—"

        if not reserva.get(
            "promocion"
        ):

            reserva["promocion"] = "No"

    cursor.close()

    return {
        "cliente": cliente,

        "total_reservas":
            total_reservas,

        "especialista_mas_atendio":
            especialista_mas_atendio,

        "servicio_mas_solicitado":
            servicio_mas_solicitado,

        "reservas":
            reservas
    }


def _generar_pdf(
    empresa,
    estadisticas,
    filtros
):

    archivo = os.path.join(
        tempfile.gettempdir(),
        f'reporte_reservas_estadistico_'
        f'{int(datetime.now().timestamp())}.pdf'
    )

    doc = SimpleDocTemplate(
        archivo,
        pagesize=LETTER,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm
    )

    doc.title = "Reporte Filtrado de Reservas"
    doc.author = "Sistema de Gestión"
    doc.subject = "Reporte Estadístico"

    styles = getSampleStyleSheet()

    titulo_style = ParagraphStyle(
        'titulo_empresa',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor("#1B365D"),
        alignment=1,
        spaceAfter=8
    )

    subtitulo_style = ParagraphStyle(
        'subtitulo',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=14,
        textColor=colors.HexColor("#1B365D"),
        alignment=1,
        spaceBefore=10,
        spaceAfter=10
    )

    normal = ParagraphStyle(
        'normal_empresa',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14
    )

    elementos = []

    logo = ""

    if empresa.get("logo"):

        base_dir = os.path.dirname(
            os.path.abspath(__file__)
        )

        logo_bd = empresa.get(
            "logo",
            ""
        ).lstrip("/")

        logo_path = os.path.abspath(
            os.path.join(
                base_dir,
                "..",
                "..",
                logo_bd
            )
        )

        print(
            "Ruta PDF:",
            logo_path
        )

        print(
            "Existe:",
            os.path.exists(logo_path)
        )

        if os.path.exists(logo_path):

            logo = Image(
                logo_path,
                width=42 * mm,
                height=42 * mm
            )

    empresa_html = f'''
    <font size="18" color="#1B365D">
    <b>
    {empresa.get(
        "nombre",
        "ESTUDIO DE BELLEZA MT"
    )}
    </b>
    </font>

    <br/><br/>

    <font size="10">

    <b>Dirección:</b>
    {empresa.get("direccion", "N/A")}

    <br/>

    <b>Teléfono:</b>
    {empresa.get("telefono", "N/A")}

    <br/>

    <b>Correo:</b>
    {empresa.get("email", "N/A")}

    </font>
    '''

    datos_empresa = Paragraph(
        empresa_html,
        normal
    )

    header = Table(
        [
            [
                logo,
                datos_empresa
            ]
        ],
        colWidths=[
            48 * mm,
            120 * mm
        ]
    )

    header.setStyle(
        TableStyle([
            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),
            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                8
            )
        ])
    )

    elementos.append(header)

    linea = Drawing(
        520,
        12
    )

    linea.add(
        Line(
            0,
            5,
            520,
            5,
            strokeColor=colors.HexColor("#D4AF37"),
            strokeWidth=2
        )
    )

    elementos.append(linea)

    elementos.append(
        Spacer(1, 12)
    )

    elementos.append(
        Paragraph(
            "REPORTE ESTADÍSTICO DE RESERVAS",
            titulo_style
        )
    )

    elementos.append(
        Paragraph(
            f'''
            Fecha de emisión:
            {datetime.now().strftime(
                '%d-%m-%Y %H:%M'
            )}
            ''',
            normal
        )
    )

    elementos.append(
        Spacer(1, 12)
    )

    elementos.append(
        Paragraph(
            "Resumen general",
            subtitulo_style
        )
    )

    metrics_data = [
        [
            "Métrica",
            "Cantidad"
        ],
        [
            "Total de reservas",
            str(
                estadisticas[
                    'total_reservas'
                ]
            )
        ],
        [
            "Reservas activas",
            str(
                estadisticas[
                    'reservas_activas'
                ]
            )
        ],
        [
            "Reservas Canceladas",
            str(
                estadisticas[
                    'reservas_inactivas'
                ]
            )
        ]
    ]

    metrics_table = Table(
        metrics_data,
        hAlign='LEFT',
        colWidths=[
            90 * mm,
            45 * mm
        ]
    )

    metrics_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (-1, 0),
                colors.HexColor("#1B365D")
            ),
            (
                'TEXTCOLOR',
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                'FONTNAME',
                (0, 0),
                (-1, 0),
                'Helvetica-Bold'
            ),
            (
                'GRID',
                (0, 0),
                (-1, -1),
                0.25,
                colors.grey
            ),
            (
                'ROWBACKGROUNDS',
                (0, 1),
                (-1, -1),
                [
                    colors.whitesmoke,
                    colors.lightgrey
                ]
            )
        ])
    )

    elementos.append(metrics_table)

    elementos.append(
        Spacer(1, 15)
    )

    def _append_section(
        titulo,
        table_data,
        col_widths
    ):

        elementos.append(
            Paragraph(
                titulo,
                subtitulo_style
            )
        )

        tabla = Table(
            table_data,
            hAlign='LEFT',
            colWidths=col_widths,
            repeatRows=1
        )

        tabla.setStyle(
            TableStyle([
                (
                    'BACKGROUND',
                    (0, 0),
                    (-1, 0),
                    colors.HexColor("#D4AF37")
                ),
                (
                    'TEXTCOLOR',
                    (0, 0),
                    (-1, 0),
                    colors.white
                ),
                (
                    'FONTNAME',
                    (0, 0),
                    (-1, 0),
                    'Helvetica-Bold'
                ),
                (
                    'GRID',
                    (0, 0),
                    (-1, -1),
                    0.25,
                    colors.grey
                ),
                (
                    'VALIGN',
                    (0, 0),
                    (-1, -1),
                    'MIDDLE'
                ),
                (
                    'FONTSIZE',
                    (0, 0),
                    (-1, -1),
                    9
                ),
                (
                    'ROWBACKGROUNDS',
                    (0, 1),
                    (-1, -1),
                    [
                        colors.whitesmoke,
                        colors.lightgrey
                    ]
                ),
                (
                    'LEFTPADDING',
                    (0, 0),
                    (-1, -1),
                    4
                ),
                (
                    'RIGHTPADDING',
                    (0, 0),
                    (-1, -1),
                    4
                ),
                (
                    'TOPPADDING',
                    (0, 0),
                    (-1, -1),
                    3
                ),
                (
                    'BOTTOMPADDING',
                    (0, 0),
                    (-1, -1),
                    3
                )
            ])
        )

        elementos.append(tabla)

        elementos.append(
            Spacer(1, 15)
        )

    if estadisticas[
        'reservas_por_servicio'
    ]:

        servicio_data = [
            [
                "Servicio",
                "Reservas"
            ]
        ]

        for item in estadisticas[
            'reservas_por_servicio'
        ]:

            servicio_data.append([
                item['servicio'],
                item['total_reservas']
            ])

        _append_section(
            "Reservas por servicio",
            servicio_data,
            [
                95 * mm,
                45 * mm
            ]
        )

    if estadisticas[
        'reservas_por_especialista'
    ]:

        especialista_data = [
            [
                "Especialista",
                "Reservas"
            ]
        ]

        for item in estadisticas[
            'reservas_por_especialista'
        ]:

            especialista_data.append([
                item['especialista'],
                item['total_reservas']
            ])

        _append_section(
            "Reservas por especialista",
            especialista_data,
            [
                95 * mm,
                45 * mm
            ]
        )

    if estadisticas[
        'top_clientes'
    ]:

        clientes_data = [
            [
                "Cliente",
                "Reservas"
            ]
        ]

        for item in estadisticas[
            'top_clientes'
        ]:

            clientes_data.append([
                item['cliente'],
                item['total_reservas']
            ])

        _append_section(
            "Clientes con mayor cantidad de reservas",
            clientes_data,
            [
                95 * mm,
                45 * mm
            ]
        )

    if estadisticas[
        'ultimas_reservas'
    ]:

        ultimas_data = [
            [
                "Cliente",
                "Especialista",
                "Fecha",
                "Hora",
                "Estado"
            ]
        ]

        for item in estadisticas[
            'ultimas_reservas'
        ]:

            estado = (
                "Activo"
                if item['status'] == 1
                else "Cancelada"
            )

            ultimas_data.append([
                item['cliente'],
                item['especialista'],
                str(
                    item['fecha_reserva']
                ),
                str(
                    item['hora_reserva']
                ),
                estado
            ])

        _append_section(
            "Últimas reservas registradas",
            ultimas_data,
            [
                45 * mm,
                40 * mm,
                25 * mm,
                20 * mm,
                20 * mm
            ]
        )

    doc.build(
        elementos,
        onFirstPage=lambda c, d: (
            borde(c, d),
            footer(c, d)
        ),
        onLaterPages=lambda c, d: (
            borde(c, d),
            footer(c, d)
        )
    )

    return archivo


def _generar_pdf_cliente(
    empresa,
    estadisticas
):

    archivo = os.path.join(
        tempfile.gettempdir(),
        f'reporte_cliente_'
        f'{int(datetime.now().timestamp())}.pdf'
    )

    doc = SimpleDocTemplate(
        archivo,
        pagesize=LETTER,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm
    )

    doc.title = "Reporte del Cliente"
    doc.author = "Sistema de Gestión"
    doc.subject = "Reporte individual del cliente"

    styles = getSampleStyleSheet()

    titulo = ParagraphStyle(
        'titulo_cliente',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor("#1B365D"),
        alignment=1,
        spaceAfter=8
    )

    subtitulo = ParagraphStyle(
        'subtitulo_cliente',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        textColor=colors.HexColor("#1B365D"),
        alignment=1,
        spaceBefore=8,
        spaceAfter=8
    )

    normal = ParagraphStyle(
        'normal_cliente',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14
    )

    elementos = []

    logo = ""

    if empresa.get("logo"):

        base_dir = os.path.dirname(
            os.path.abspath(__file__)
        )

        logo_bd = empresa.get(
            "logo",
            ""
        ).lstrip("/")

        logo_path = os.path.abspath(
            os.path.join(
                base_dir,
                "..",
                "..",
                logo_bd
            )
        )

        if os.path.exists(logo_path):

            logo = Image(
                logo_path,
                width=35 * mm,
                height=35 * mm
            )

    empresa_html = f'''
    <font size="18" color="#1B365D">
    <b>
    {empresa.get(
        "nombre",
        "ESTUDIO DE BELLEZA MT"
    )}
    </b>
    </font>

    <br/><br/>

    <font size="9">

    <b>Dirección:</b>
    {empresa.get("direccion", "N/A")}

    <br/>

    <b>Teléfono:</b>
    {empresa.get("telefono", "N/A")}

    <br/>

    <b>Correo:</b>
    {empresa.get("email", "N/A")}

    </font>
    '''

    datos_empresa = Paragraph(
        empresa_html,
        normal
    )

    header = Table(
        [
            [
                logo,
                datos_empresa
            ]
        ],
        colWidths=[
            45 * mm,
            135 * mm
        ]
    )

    header.setStyle(
        TableStyle([
            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),
            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                6
            )
        ])
    )

    elementos.append(header)

    linea = Drawing(
        520,
        12
    )

    linea.add(
        Line(
            0,
            5,
            520,
            5,
            strokeColor=colors.HexColor("#D4AF37"),
            strokeWidth=2
        )
    )

    elementos.append(linea)

    elementos.append(
        Spacer(1, 8)
    )

    elementos.append(
        Paragraph(
            "REPORTE DEL CLIENTE",
            titulo
        )
    )

    cliente = estadisticas.get(
        "cliente",
        {}
    )

    nombre_cliente = (
        f"{cliente.get('nombre_usuario', '')} "
        f"{cliente.get('apellido_usuario', '')}"
    ).strip()

    elementos.append(
        Paragraph(
            f'''
            <b>Cliente:</b>
            {nombre_cliente.upper()}

            <br/>

            <b>Cédula:</b>
            {cliente.get('cedula', 'N/A')}

            <br/>

            <b>Fecha del reporte:</b>
            {datetime.now().strftime('%d/%m/%Y')}
            ''',
            normal
        )
    )

    elementos.append(
        Spacer(1, 12)
    )

    elementos.append(
        Paragraph(
            "RESUMEN DEL CLIENTE",
            subtitulo
        )
    )

    total_reservas = estadisticas.get(
        "total_reservas",
        0
    )

    especialista = estadisticas.get(
        "especialista_mas_atendio",
        "—"
    )

    servicio = estadisticas.get(
        "servicio_mas_solicitado",
        "—"
    )

    resumen_data = [
        [
            "TOTAL RESERVAS",
            "ESPECIALISTA QUE MÁS LA ATENDIÓ",
            "SERVICIO MÁS SOLICITADO"
        ],
        [
            str(total_reservas),
            especialista,
            servicio
        ]
    ]

    resumen_table = Table(
        resumen_data,
        colWidths=[
            55 * mm,
            65 * mm,
            55 * mm
        ],
        rowHeights=[
            12 * mm,
            20 * mm
        ]
    )

    resumen_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (-1, 0),
                colors.HexColor("#1B365D")
            ),
            (
                'TEXTCOLOR',
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                'FONTNAME',
                (0, 0),
                (-1, 0),
                'Helvetica-Bold'
            ),
            (
                'FONTSIZE',
                (0, 0),
                (-1, 0),
                8
            ),
            (
                'FONTNAME',
                (0, 1),
                (-1, 1),
                'Helvetica-Bold'
            ),
            (
                'FONTSIZE',
                (0, 1),
                (-1, 1),
                9
            ),
            (
                'TEXTCOLOR',
                (0, 1),
                (-1, 1),
                colors.HexColor("#1B365D")
            ),
            (
                'BACKGROUND',
                (0, 1),
                (-1, 1),
                colors.whitesmoke
            ),
            (
                'GRID',
                (0, 0),
                (-1, -1),
                0.5,
                colors.HexColor("#D4AF37")
            ),
            (
                'ALIGN',
                (0, 0),
                (-1, -1),
                'CENTER'
            ),
            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),
            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                5
            ),
            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                5
            )
        ])
    )

    elementos.append(
        resumen_table
    )

    elementos.append(
        Spacer(1, 15)
    )

    elementos.append(
        Paragraph(
            "HISTORIAL DE RESERVAS",
            subtitulo
        )
    )

    reservas = estadisticas.get(
        "reservas",
        []
    )

    historial_data = [
        [
            "Fecha",
            "Hora",
            "Especialista",
            "Servicio",
            "Promoción",
            "Estado"
        ]
    ]

    for reserva in reservas:

        historial_data.append([
            reserva.get(
                "fecha",
                "—"
            ),
            reserva.get(
                "hora",
                "—"
            ),
            reserva.get(
                "especialista",
                "—"
            ),
            reserva.get(
                "servicio",
                "—"
            ),
            reserva.get(
                "promocion",
                "No"
            ),
            reserva.get(
                "estado",
                "—"
            )
        ])

    if len(historial_data) == 1:

        historial_data.append([
            "—",
            "—",
            "—",
            "—",
            "—",
            "Sin reservas"
        ])

    historial_table = Table(
        historial_data,
        colWidths=[
            25 * mm,
            23 * mm,
            48 * mm,
            42 * mm,
            22 * mm,
            25 * mm
        ],
        repeatRows=1
    )

    historial_table.setStyle(
        TableStyle([
            (
                'BACKGROUND',
                (0, 0),
                (-1, 0),
                colors.HexColor("#D4AF37")
            ),
            (
                'TEXTCOLOR',
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                'FONTNAME',
                (0, 0),
                (-1, 0),
                'Helvetica-Bold'
            ),
            (
                'FONTSIZE',
                (0, 0),
                (-1, 0),
                8
            ),
            (
                'FONTSIZE',
                (0, 1),
                (-1, -1),
                7.5
            ),
            (
                'GRID',
                (0, 0),
                (-1, -1),
                0.4,
                colors.grey
            ),
            (
                'VALIGN',
                (0, 0),
                (-1, -1),
                'MIDDLE'
            ),
            (
                'ALIGN',
                (0, 0),
                (-1, -1),
                'CENTER'
            ),
            (
                'ROWBACKGROUNDS',
                (0, 1),
                (-1, -1),
                [
                    colors.whitesmoke,
                    colors.lightgrey
                ]
            ),
            (
                'LEFTPADDING',
                (0, 0),
                (-1, -1),
                3
            ),
            (
                'RIGHTPADDING',
                (0, 0),
                (-1, -1),
                3
            ),
            (
                'TOPPADDING',
                (0, 0),
                (-1, -1),
                4
            ),
            (
                'BOTTOMPADDING',
                (0, 0),
                (-1, -1),
                4
            )
        ])
    )

    elementos.append(
        historial_table
    )

    doc.build(
        elementos,
        onFirstPage=lambda c, d: (
            borde(c, d),
            footer(c, d)
        ),
        onLaterPages=lambda c, d: (
            borde(c, d),
            footer(c, d)
        )
    )

    return archivo


# ==============================================================
# EXCEL GENERAL
# ==============================================================

def _generar_excel(
    empresa,
    estadisticas,
    filtros
):

    workbook = Workbook()

    hoja = workbook.active

    hoja.title = "Reporte de Reservas"

    color_azul = "1B365D"
    color_dorado = "D4AF37"
    color_gris = "F2F2F2"
    color_gris_oscuro = "666666"
    color_blanco = "FFFFFF"
    color_borde = "D9D9D9"

    thin = Side(
        border_style="thin",
        color=color_borde
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    hoja.sheet_view.showGridLines = False

    for columna in [
        "A",
        "B",
        "C",
        "D",
        "E",
        "F"
    ]:

        hoja.column_dimensions[
            columna
        ].width = 24

    hoja.column_dimensions[
        "A"
    ].width = 40

    hoja.merge_cells(
        "A1:F1"
    )

    hoja["A1"] = empresa.get(
        "nombre_empresa",
        "ESTUDIO DE BELLEZA"
    )

    hoja["A1"].font = Font(
        size=20,
        bold=True,
        color=color_blanco
    )

    hoja["A1"].fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    hoja["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        1
    ].height = 35

    hoja.merge_cells(
        "A2:F2"
    )

    hoja["A2"] = (
        "REPORTE ESTADÍSTICO DE RESERVAS"
    )

    hoja["A2"].font = Font(
        size=15,
        bold=True,
        color=color_dorado
    )

    hoja["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        2
    ].height = 28

    hoja.merge_cells(
        "A3:F3"
    )

    hoja["A3"] = (
        "Resumen general del comportamiento "
        "de las reservas"
    )

    hoja["A3"].font = Font(
        size=10,
        italic=True,
        color=color_gris_oscuro
    )

    hoja["A3"].alignment = Alignment(
        horizontal="center"
    )

    hoja.merge_cells(
        "A5:F5"
    )

    hoja["A5"] = (
        f"Fecha de generación: "
        f"{datetime.now().strftime('%d/%m/%Y %H:%M')}"
    )

    hoja["A5"].font = Font(
        size=10,
        color=color_gris_oscuro
    )

    hoja["A5"].alignment = Alignment(
        horizontal="right"
    )

    fila = 7

    hoja.merge_cells(
        start_row=fila,
        start_column=1,
        end_row=fila,
        end_column=6
    )

    celda = hoja.cell(
        fila,
        1
    )

    celda.value = "RESUMEN GENERAL"

    celda.font = Font(
        bold=True,
        size=13,
        color=color_blanco
    )

    celda.fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    celda.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        fila
    ].height = 25

    fila += 2

    resumen = [
        (
            "Total de reservas",
            estadisticas.get(
                "total_reservas",
                0
            )
        ),
        (
            "Reservas activas",
            estadisticas.get(
                "reservas_activas",
                0
            )
        ),
        (
            "Reservas inactivas",
            estadisticas.get(
                "reservas_inactivas",
                0
            )
        )
    ]

    for columna, (
        titulo,
        valor
    ) in enumerate(
        resumen,
        1
    ):

        celda_titulo = hoja.cell(
            fila,
            columna
        )

        celda_valor = hoja.cell(
            fila + 1,
            columna
        )

        celda_titulo.value = titulo

        celda_valor.value = valor

        celda_titulo.font = Font(
            bold=True,
            color=color_blanco
        )

        celda_titulo.fill = PatternFill(
            "solid",
            fgColor=color_azul
        )

        celda_titulo.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        celda_titulo.border = border

        celda_valor.font = Font(
            bold=True,
            size=17,
            color=color_azul
        )

        celda_valor.fill = PatternFill(
            "solid",
            fgColor=color_gris
        )

        celda_valor.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        celda_valor.border = border

    fila += 4

    hoja.merge_cells(
        start_row=fila,
        start_column=1,
        end_row=fila,
        end_column=3
    )

    celda = hoja.cell(
        fila,
        1
    )

    celda.value = "RESERVAS POR SERVICIO"

    celda.font = Font(
        bold=True,
        size=12,
        color=color_blanco
    )

    celda.fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    celda.alignment = Alignment(
        horizontal="center"
    )

    fila += 1

    encabezados = [
        "Servicio",
        "Cantidad",
        "Porcentaje"
    ]

    for columna, titulo in enumerate(
        encabezados,
        1
    ):

        celda = hoja.cell(
            fila,
            columna
        )

        celda.value = titulo

        celda.font = Font(
            bold=True,
            color=color_azul
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=color_gris
        )

        celda.border = border

        celda.alignment = Alignment(
            horizontal="center"
        )

    fila += 1

    servicios = estadisticas.get(
        "reservas_por_servicio",
        []
    )

    total_reservas = estadisticas.get(
        "total_reservas",
        0
    )

    for servicio in servicios:

        nombre = (
            servicio.get(
                "nombre_servicio"
            )
            or servicio.get(
                "servicio"
            )
            or "Sin nombre"
        )

        cantidad = (
            servicio.get(
                "cantidad"
            )
            or servicio.get(
                "total"
            )
            or servicio.get(
                "total_reservas"
            )
            or 0
        )

        porcentaje = (
            cantidad / total_reservas
            if total_reservas > 0
            else 0
        )

        hoja.cell(
            fila,
            1
        ).value = nombre

        hoja.cell(
            fila,
            2
        ).value = cantidad

        hoja.cell(
            fila,
            3
        ).value = porcentaje

        hoja.cell(
            fila,
            3
        ).number_format = "0.00%"

        for columna in range(
            1,
            4
        ):

            hoja.cell(
                fila,
                columna
            ).border = border

        fila += 1

    fila += 2

    hoja.merge_cells(
        start_row=fila,
        start_column=1,
        end_row=fila,
        end_column=3
    )

    celda = hoja.cell(
        fila,
        1
    )

    celda.value = "RESERVAS POR ESPECIALISTA"

    celda.font = Font(
        bold=True,
        size=12,
        color=color_blanco
    )

    celda.fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    celda.alignment = Alignment(
        horizontal="center"
    )

    fila += 1

    encabezados = [
        "Especialista",
        "Cantidad",
        "Porcentaje"
    ]

    for columna, titulo in enumerate(
        encabezados,
        1
    ):

        celda = hoja.cell(
            fila,
            columna
        )

        celda.value = titulo

        celda.font = Font(
            bold=True,
            color=color_azul
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=color_gris
        )

        celda.border = border

        celda.alignment = Alignment(
            horizontal="center"
        )

    fila += 1

    especialistas = estadisticas.get(
        "reservas_por_especialista",
        []
    )

    for especialista in especialistas:

        nombre = (
            especialista.get(
                "nombre_especialista"
            )
            or especialista.get(
                "especialista"
            )
            or "Sin nombre"
        )

        cantidad = (
            especialista.get(
                "cantidad"
            )
            or especialista.get(
                "total"
            )
            or especialista.get(
                "total_reservas"
            )
            or 0
        )

        porcentaje = (
            cantidad / total_reservas
            if total_reservas > 0
            else 0
        )

        hoja.cell(
            fila,
            1
        ).value = nombre

        hoja.cell(
            fila,
            2
        ).value = cantidad

        hoja.cell(
            fila,
            3
        ).value = porcentaje

        hoja.cell(
            fila,
            3
        ).number_format = "0.00%"

        for columna in range(
            1,
            4
        ):

            hoja.cell(
                fila,
                columna
            ).border = border

        fila += 1

    fila += 2

    hoja.merge_cells(
        start_row=fila,
        start_column=1,
        end_row=fila,
        end_column=3
    )

    celda = hoja.cell(
        fila,
        1
    )

    celda.value = "CLIENTES CON MÁS RESERVAS"

    celda.font = Font(
        bold=True,
        size=12,
        color=color_blanco
    )

    celda.fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    celda.alignment = Alignment(
        horizontal="center"
    )

    fila += 1

    encabezados = [
        "Cliente",
        "Reservas",
        "Porcentaje"
    ]

    for columna, titulo in enumerate(
        encabezados,
        1
    ):

        celda = hoja.cell(
            fila,
            columna
        )

        celda.value = titulo

        celda.font = Font(
            bold=True,
            color=color_azul
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=color_gris
        )

        celda.border = border

        celda.alignment = Alignment(
            horizontal="center"
        )

    fila += 1

    clientes = estadisticas.get(
        "top_clientes",
        []
    )

    for cliente in clientes:

        nombre = (
            cliente.get(
                "nombre_cliente"
            )
            or cliente.get(
                "cliente"
            )
            or "Sin nombre"
        )

        cantidad = (
            cliente.get(
                "cantidad"
            )
            or cliente.get(
                "total"
            )
            or cliente.get(
                "total_reservas"
            )
            or 0
        )

        porcentaje = (
            cantidad / total_reservas
            if total_reservas > 0
            else 0
        )

        hoja.cell(
            fila,
            1
        ).value = nombre

        hoja.cell(
            fila,
            2
        ).value = cantidad

        hoja.cell(
            fila,
            3
        ).value = porcentaje

        hoja.cell(
            fila,
            3
        ).number_format = "0.00%"

        for columna in range(
            1,
            4
        ):

            hoja.cell(
                fila,
                columna
            ).border = border

        fila += 1

    for fila_excel in hoja.iter_rows():

        for celda in fila_excel:

            celda.alignment = Alignment(
                vertical="center",
                wrap_text=True
            )

    hoja.freeze_panes = "A8"

    hoja.page_setup.orientation = "landscape"

    hoja.page_setup.fitToWidth = 1

    hoja.page_setup.fitToHeight = 0

    hoja.sheet_properties.pageSetUpPr.fitToPage = True

    hoja.page_margins.left = 0.25

    hoja.page_margins.right = 0.25

    hoja.page_margins.top = 0.5

    hoja.page_margins.bottom = 0.5

    nombre_archivo = (
        "reporte_estadistico_reservas_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    ruta_archivo = os.path.join(
        "reportes",
        nombre_archivo
    )

    os.makedirs(
        "reportes",
        exist_ok=True
    )

    workbook.save(
        ruta_archivo
    )

    return ruta_archivo


# ==============================================================
# EXCEL INDIVIDUAL DEL CLIENTE
# ==============================================================

def _generar_excel_cliente(
    empresa,
    estadisticas
):

    workbook = Workbook()

    hoja = workbook.active

    hoja.title = "Reporte del Cliente"

    # ==========================================================
    # COLORES
    # ==========================================================

    color_azul = "1B365D"
    color_dorado = "D4AF37"
    color_gris = "F2F2F2"
    color_gris_oscuro = "666666"
    color_blanco = "FFFFFF"
    color_borde = "D9D9D9"

    thin = Side(
        border_style="thin",
        color=color_borde
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )

    hoja.sheet_view.showGridLines = False

    # ==========================================================
    # ANCHO DE COLUMNAS
    # ==========================================================

    hoja.column_dimensions["A"].width = 15
    hoja.column_dimensions["B"].width = 15
    hoja.column_dimensions["C"].width = 38
    hoja.column_dimensions["D"].width = 40
    hoja.column_dimensions["E"].width = 15
    hoja.column_dimensions["F"].width = 15

    # ==========================================================
    # DATOS DEL CLIENTE
    # ==========================================================

    cliente = estadisticas.get(
        "cliente",
        {}
    )

    nombre_cliente = (
        f"{cliente.get('nombre_usuario', '')} "
        f"{cliente.get('apellido_usuario', '')}"
    ).strip()

    cedula_cliente = cliente.get(
        "cedula",
        "N/A"
    )

    nombre_empresa = empresa.get(
        "nombre",
        empresa.get(
            "nombre_empresa",
            "ESTUDIO DE BELLEZA "
        )
    )

    # ==========================================================
    # ENCABEZADO
    # ==========================================================

    hoja.merge_cells(
        "A1:F1"
    )

    hoja["A1"] = nombre_empresa

    hoja["A1"].font = Font(
        size=20,
        bold=True,
        color=color_blanco
    )

    hoja["A1"].fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    hoja["A1"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        1
    ].height = 35

    # ==========================================================
    # TÍTULO
    # ==========================================================

    hoja.merge_cells(
        "A2:F2"
    )

    hoja["A2"] = (
        "REPORTE DEL CLIENTE"
    )

    hoja["A2"].font = Font(
        size=15,
        bold=True,
        color=color_dorado
    )

    hoja["A2"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        2
    ].height = 28

    # ==========================================================
    # INFORMACIÓN DEL CLIENTE
    # ==========================================================

    hoja.merge_cells(
        "A4:F4"
    )

    hoja["A4"] = (
        "INFORMACIÓN DEL CLIENTE"
    )

    hoja["A4"].font = Font(
        bold=True,
        size=12,
        color=color_blanco
    )

    hoja["A4"].fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    hoja["A4"].alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        4
    ].height = 25

    hoja["A5"] = "Cliente"

    hoja["B5"] = (
        nombre_cliente.upper()
    )

    hoja["A6"] = "Cédula"

    hoja["B6"] = cedula_cliente

    hoja["A7"] = "Fecha del reporte"

    hoja["B7"] = datetime.now().strftime(
        "%d/%m/%Y"
    )

    for fila in range(
        5,
        8
    ):

        hoja.cell(
            fila,
            1
        ).font = Font(
            bold=True,
            color=color_azul
        )

        hoja.cell(
            fila,
            1
        ).fill = PatternFill(
            "solid",
            fgColor=color_gris
        )

        hoja.cell(
            fila,
            1
        ).border = border

        hoja.cell(
            fila,
            2
        ).border = border

        hoja.cell(
            fila,
            2
        ).alignment = Alignment(
            vertical="center",
            wrap_text=True
        )

    # ==========================================================
    # RESUMEN DEL CLIENTE
    # ==========================================================

    fila = 10

    hoja.merge_cells(
        start_row=fila,
        start_column=1,
        end_row=fila,
        end_column=6
    )

    celda = hoja.cell(
        fila,
        1
    )

    celda.value = (
        "RESUMEN DEL CLIENTE"
    )

    celda.font = Font(
        bold=True,
        size=12,
        color=color_blanco
    )

    celda.fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    celda.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        fila
    ].height = 25

    fila += 1

    # ==========================================================
    # MÉTRICAS
    # ==========================================================

    total_reservas = estadisticas.get(
        "total_reservas",
        0
    )

    especialista = estadisticas.get(
        "especialista_mas_atendio",
        "—"
    )

    servicio = estadisticas.get(
        "servicio_mas_solicitado",
        "—"
    )

    metricas = [
        (
            "TOTAL RESERVAS",
            total_reservas
        ),
        (
            "ESPECIALISTA QUE MÁS LA ATENDIÓ",
            especialista
        ),
        (
            "SERVICIO MÁS SOLICITADO",
            servicio
        )
    ]

    columnas_metricas = [
        (1, 2),
        (3, 4),
        (5, 6)
    ]

    for indice, (
        titulo,
        valor
    ) in enumerate(metricas):

        inicio, fin = columnas_metricas[
            indice
        ]

        hoja.merge_cells(
            start_row=fila,
            start_column=inicio,
            end_row=fila,
            end_column=fin
        )

        celda = hoja.cell(
            fila,
            inicio
        )

        celda.value = titulo

        celda.font = Font(
            bold=True,
            size=9,
            color=color_blanco
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=color_azul
        )

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        for columna in range(
            inicio,
            fin + 1
        ):

            hoja.cell(
                fila,
                columna
            ).border = border

    fila += 1

    for indice, (
        titulo,
        valor
    ) in enumerate(metricas):

        inicio, fin = columnas_metricas[
            indice
        ]

        hoja.merge_cells(
            start_row=fila,
            start_column=inicio,
            end_row=fila,
            end_column=fin
        )

        celda = hoja.cell(
            fila,
            inicio
        )

        celda.value = valor

        celda.font = Font(
            bold=True,
            size=11,
            color=color_azul
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=color_gris
        )

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

        for columna in range(
            inicio,
            fin + 1
        ):

            hoja.cell(
                fila,
                columna
            ).border = border

    hoja.row_dimensions[
        fila
    ].height = 40

    # ==========================================================
    # HISTORIAL DE RESERVAS
    # ==========================================================

    fila += 3

    hoja.merge_cells(
        start_row=fila,
        start_column=1,
        end_row=fila,
        end_column=6
    )

    celda = hoja.cell(
        fila,
        1
    )

    celda.value = (
        "HISTORIAL DE RESERVAS"
    )

    celda.font = Font(
        bold=True,
        size=12,
        color=color_blanco
    )

    celda.fill = PatternFill(
        "solid",
        fgColor=color_azul
    )

    celda.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    hoja.row_dimensions[
        fila
    ].height = 25

    fila += 1

    encabezados = [
        "Fecha",
        "Hora",
        "Especialista",
        "Servicio",
        "Promoción",
        "Estado"
    ]

    for columna, titulo in enumerate(
        encabezados,
        1
    ):

        celda = hoja.cell(
            fila,
            columna
        )

        celda.value = titulo

        celda.font = Font(
            bold=True,
            color=color_blanco
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=color_dorado
        )

        celda.border = border

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )

    fila += 1


    reservas = estadisticas.get(
        "reservas",
        []
    )

    for reserva in reservas:

        fecha = reserva.get(
            "fecha",
            "—"
        )

        hora = reserva.get(
            "hora",
            "—"
        )

        especialista_reserva = reserva.get(
            "especialista",
            "—"
        )

        servicio_reserva = reserva.get(
            "servicio",
            "—"
        )

        promocion = reserva.get(
            "promocion",
            "No"
        )

        estado = reserva.get(
            "estado",
            "—"
        )

        hoja.cell(
            fila,
            1
        ).value = fecha

        hoja.cell(
            fila,
            2
        ).value = hora

        hoja.cell(
            fila,
            3
        ).value = especialista_reserva

        hoja.cell(
            fila,
            4
        ).value = servicio_reserva

        hoja.cell(
            fila,
            5
        ).value = promocion

        hoja.cell(
            fila,
            6
        ).value = estado

        for columna in range(
            1,
            7
        ):

            celda = hoja.cell(
                fila,
                columna
            )

            celda.border = border

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )

        fila += 1

    # ==========================================================
    # SIN RESERVAS
    # ==========================================================

    if not reservas:

        hoja.merge_cells(
            start_row=fila,
            start_column=1,
            end_row=fila,
            end_column=6
        )

        celda = hoja.cell(
            fila,
            1
        )

        celda.value = (
            "Este cliente no posee reservas registradas."
        )

        celda.font = Font(
            italic=True,
            color=color_gris_oscuro
        )

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        for columna in range(
            1,
            7
        ):

            hoja.cell(
                fila,
                columna
            ).border = border

    # ==========================================================
    # CONFIGURACIÓN DE IMPRESIÓN
    # ==========================================================

    hoja.freeze_panes = "A14"

    hoja.page_setup.orientation = "landscape"

    hoja.page_setup.fitToWidth = 1

    hoja.page_setup.fitToHeight = 0

    hoja.sheet_properties.pageSetUpPr.fitToPage = True

    hoja.page_margins.left = 0.25

    hoja.page_margins.right = 0.25

    hoja.page_margins.top = 0.5

    hoja.page_margins.bottom = 0.5

    # ==========================================================
    # ARCHIVO
    # ==========================================================

    nombre_archivo = (
        "reporte_cliente_"
        f"{cedula_cliente}_"
        f"{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    ruta_archivo = os.path.join(
        "reportes",
        nombre_archivo
    )

    os.makedirs(
        "reportes",
        exist_ok=True
    )

    workbook.save(
        ruta_archivo
    )

    return ruta_archivo


# ==============================================================
# HANDLE REQUEST
# ==============================================================

def handle_request(ruta):

    modelo = ReservasModelo()

    clientes = modelo.consultar_clientes()
    especialistasr = modelo.consultar_especialistasreporte()

    # ==========================================================
    # REPORTE PDF INDIVIDUAL DEL CLIENTE
    # ==========================================================

    if (
        request.method == "POST"
        and request.form.get("pdf_cliente")
    ):

        cedula_cliente = request.form.get(
            "cliente_reporte"
        )

        if not cedula_cliente:

            return (
                "Debe seleccionar un cliente",
                400
            )

        empresa = EmpresaModelo().mostrar()

        empresa_data = (
            empresa[0]
            if empresa
            else {}
        )

        estadisticas_cliente = (
            _consultar_estadisticas_cliente(
                modelo,
                cedula_cliente
            )
        )

        archivo = _generar_pdf_cliente(
            empresa_data,
            estadisticas_cliente
        )

        return send_file(
            archivo,
            mimetype="application/pdf",
            as_attachment=False
        )


    # ==========================================================
    # REPORTE GENERAL
    # ==========================================================

    panel = PanelModelo()

    reservas_mes = (
        panel.reservas_por_mes()
    )

    reservas_dia = (
        panel.reservas_por_dia()
    )

    reservas_hora = (
        panel.reservas_por_hora()
    )

    desde = (
        request.form.get("desde")
        or request.args.get("desde")
    )

    hasta = (
        request.form.get("hasta")
        or request.args.get("hasta")
    )

    status = request.form.get(
        "status",
        "all"
    )

    filtros = {
        "desde": desde or "",
        "hasta": hasta or "",
        "status": status
    }

    estadisticas = _consultar_estadisticas(
        modelo,
        desde=desde,
        hasta=hasta,
        status=status
    )

    # ==========================================================
    # PDF GENERAL
    # ==========================================================

    if (
        request.method == "POST"
        and request.form.get("pdf")
    ):

        empresa = EmpresaModelo().mostrar()

        empresa_data = (
            empresa[0]
            if empresa
            else {}
        )

        archivo = _generar_pdf(
            empresa_data,
            estadisticas,
            filtros
        )

        return send_file(
            archivo,
            mimetype="application/pdf",
            as_attachment=False
        )

    # ==========================================================
    # EXCEL GENERAL
    # ==========================================================

    if (
        request.method == "POST"
        and request.form.get("excel")
    ):

        empresa = EmpresaModelo().mostrar()

        empresa_data = (
            empresa[0]
            if empresa
            else {}
        )

        archivo = _generar_excel(
            empresa_data,
            estadisticas,
            filtros
        )

        return send_file(
            archivo,
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            as_attachment=True,
            download_name=(
                "reporte_reservas_estadistico.xlsx"
            )
        )

    # ==========================================================
    # RESPUESTA NORMAL
    # ==========================================================

    return {
        "reservas_mes": reservas_mes,
        "reservas_dia": reservas_dia,
        "reservas_hora": reservas_hora,
        "estadisticas": estadisticas,
        "especialistasr": especialistasr,
        "clientes": clientes,
        "filtros": filtros
    }