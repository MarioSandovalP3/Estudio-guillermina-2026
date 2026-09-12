from flask import request, send_file
from admin.modelo.pagos import PagosModelo
from admin.modelo.empresa import EmpresaModelo
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, Image
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Drawing, Line
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
import tempfile
import os


def _logo_path():
    return os.path.abspath(os.path.join(
        os.path.dirname(os.path.abspath(__file__)), '..', '..', 'static', 'Logo SAC.jpeg'
    ))


def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('Helvetica', 9)
    canvas.setFillColor(colors.grey)
    canvas.drawString(15 * mm, 12 * mm, 'Sistema de Gestión | Reporte generado automáticamente')
    canvas.drawRightString(195 * mm, 12 * mm, f'Página {doc.page}')
    canvas.restoreState()


def borde(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.lightgrey)
    canvas.setLineWidth(1)
    canvas.rect(20 * mm, 20 * mm, LETTER[0] - 40 * mm, LETTER[1] - 40 * mm)
    canvas.restoreState()


def _build_filters(desde, hasta, cod_tipo_pago, status):
    condiciones = []
    parametros = []

    if desde:
        condiciones.append('p.fecha_pago >= %s')
        parametros.append(desde)

    if hasta:
        condiciones.append('p.fecha_pago <= %s')
        parametros.append(hasta)

    if cod_tipo_pago and cod_tipo_pago > 0:
        condiciones.append('dp.cod_tipo_pago = %s')
        parametros.append(cod_tipo_pago)

    if status == 'active':
        condiciones.append('p.status = 1')
    elif status == 'anulado':
        condiciones.append('p.status = 0')

    where = ' AND '.join(condiciones) if condiciones else '1=1'
    return where, parametros


def _fetch_scalar(modelo, sql, parametros):
    cursor = modelo.conexion.cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    fila = cursor.fetchone()
    cursor.close()
    return fila


def _fetch_rows(modelo, sql, parametros):
    cursor = modelo.conexion.cursor(dictionary=True)
    cursor.execute(sql, tuple(parametros))
    filas = cursor.fetchall()
    cursor.close()
    return filas


def _consultar_estadisticas(modelo, desde=None, hasta=None, cod_tipo_pago=0, status='all'):
    where, parametros = _build_filters(desde, hasta, cod_tipo_pago, status)

    total_pagos = _fetch_scalar(modelo, f'''
        SELECT COUNT(DISTINCT p.cod_pago) AS total_pagos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
    ''', parametros) or {}

    ingresos = _fetch_scalar(modelo, f'''
        SELECT
            COALESCE(SUM(p.monto), 0) AS ingresos_totales,
            COALESCE(AVG(p.monto), 0) AS pago_promedio
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
    ''', parametros) or {}

    estado = _fetch_scalar(modelo, f'''
        SELECT
            SUM(p.status = 1) AS pagos_activos,
            SUM(p.status = 0) AS pagos_anulados
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
    ''', parametros) or {}

    pagos_por_tipo = _fetch_rows(modelo, f'''
        SELECT
            COALESCE(tp.cod_tipo_pago, 0) AS cod_tipo_pago,
            COALESCE(tp.nombre_tipo_pago, 'Sin tipo') AS nombre_tipo_pago,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN tipo_pago tp ON tp.cod_tipo_pago = dp.cod_tipo_pago
        WHERE {where}
        GROUP BY tp.cod_tipo_pago, tp.nombre_tipo_pago
        ORDER BY ingresos DESC
    ''', parametros)

    pagos_por_moneda = _fetch_rows(modelo, f'''
        SELECT
            COALESCE(m.cod_moneda, 0) AS cod_moneda,
            COALESCE(m.nombre_moneda, 'Bs') AS nombre_moneda,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN moneda m ON m.cod_moneda = COALESCE(dp.cod_moneda, 1)
        WHERE {where}
        GROUP BY m.cod_moneda, m.nombre_moneda
        ORDER BY ingresos DESC
    ''', parametros)

    pagos_por_mes = _fetch_rows(modelo, f'''
        SELECT
            DATE_FORMAT(p.fecha_pago, '%%b %%Y') AS periodo,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
        GROUP BY YEAR(p.fecha_pago), MONTH(p.fecha_pago)
        ORDER BY YEAR(p.fecha_pago), MONTH(p.fecha_pago)
        LIMIT 12
    ''', parametros)

    top_clientes = _fetch_rows(modelo, f'''
        SELECT
            u.cedula,
            u.nombre_usuario AS nombre_cliente,
            COUNT(DISTINCT p.cod_pago) AS total_pagos,
            COALESCE(SUM(p.monto), 0) AS ingresos
        FROM pago p
        INNER JOIN servicio s ON p.cod_servicio = s.cod_servicio
        INNER JOIN seguridad.usuario u ON s.cedula = u.cedula
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        WHERE {where}
        GROUP BY u.cedula, u.nombre_usuario
        ORDER BY ingresos DESC
        LIMIT 5
    ''', parametros)

    ultimos_pagos = _fetch_rows(modelo, f'''
        SELECT
            p.cod_pago,
            p.fecha_pago,
            p.monto,
            p.status,
            u.nombre_usuario AS cliente,
            COALESCE(tp.nombre_tipo_pago, 'N/A') AS tipo_pago,
            COALESCE(m.nombre_moneda, 'Bs') AS moneda
        FROM pago p
        INNER JOIN servicio s ON p.cod_servicio = s.cod_servicio
        INNER JOIN seguridad.usuario u ON s.cedula = u.cedula
        LEFT JOIN detalle_pago dp ON dp.cod_pago = p.cod_pago
        LEFT JOIN tipo_pago tp ON tp.cod_tipo_pago = dp.cod_tipo_pago
        LEFT JOIN moneda m ON m.cod_moneda = COALESCE(dp.cod_moneda, 1)
        WHERE {where}
        ORDER BY p.fecha_pago DESC
        LIMIT 10
    ''', parametros)

    return {
        'total_pagos': int(total_pagos.get('total_pagos', 0)),
        'ingresos_totales': float(ingresos.get('ingresos_totales', 0.0)),
        'pago_promedio': float(ingresos.get('pago_promedio', 0.0)),
        'pagos_activos': int(estado.get('pagos_activos', 0)),
        'pagos_anulados': int(estado.get('pagos_anulados', 0)),
        'pagos_por_tipo': pagos_por_tipo,
        'pagos_por_moneda': pagos_por_moneda,
        'pagos_por_mes': pagos_por_mes,
        'top_clientes': top_clientes,
        'ultimos_pagos': ultimos_pagos,
    }


def _generar_pdf(empresa, estadisticas, filtros):
    archivo = os.path.join(tempfile.gettempdir(), f'reporte_pagos_estadistico_{int(datetime.now().timestamp())}.pdf')
    doc = SimpleDocTemplate(
        archivo,
        pagesize=LETTER,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=25 * mm,
        bottomMargin=25 * mm
    )
    styles = getSampleStyleSheet()

    empresa_style = ParagraphStyle(
        'empresa', parent=styles['Heading1'], fontName='Helvetica-Bold',
        fontSize=18, leading=20, textColor=colors.HexColor('#1B365D'), alignment=TA_LEFT
    )
    titulo_style = ParagraphStyle(
        'titulo', parent=styles['Heading1'], fontName='Helvetica-Bold',
        fontSize=18, textColor=colors.HexColor('#B8860B'), alignment=TA_CENTER, spaceAfter=4
    )
    subtitulo_style = ParagraphStyle(
        'subtitulo', parent=styles['Normal'], fontName='Helvetica', fontSize=10,
        textColor=colors.grey, alignment=TA_CENTER
    )

    normal_style = ParagraphStyle(
        "normal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14
    )

    elementos = []

    logo = ''
    logo_path = _logo_path()
    if os.path.exists(logo_path):
        logo = Image(logo_path, width=42 * mm, height=42 * mm)

    direccion_empresa = empresa.get('direccion') or 'Carrera 13 entre calles 55 y 56'
    telefono_empresa = empresa.get('telefono') or '04145513219'
    correo_empresa = empresa.get('email') or 'guilermina@gmail.com'

    empresa_html = f"""
        <font size="18" color="#1B365D"><b>ESTUDIO DE BELLEZA GUILLERMINA</b></font><br/><br/>
        <font size="10"><b>Dirección:</b> {direccion_empresa}<br/>
        <b>Teléfono:</b> {telefono_empresa}<br/>
        <b>Correo:</b> {correo_empresa}</font>
    """
    encabezado = Table([[logo, Paragraph(empresa_html, normal_style)]], colWidths=[48 * mm, 120 * mm])
    encabezado.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8)
    ]))
    elementos.append(encabezado)

    linea = Drawing(520, 12)
    linea.add(Line(0, 5, 520, 5, strokeColor=colors.HexColor('#D4AF37'), strokeWidth=2))
    elementos.extend([linea, Spacer(1, 12)])
    elementos.append(Paragraph('REPORTE ESTADÍSTICO DE PAGOS', titulo_style))
    elementos.append(Paragraph('Análisis de pagos registrados en el sistema administrativo.', subtitulo_style))
    elementos.append(Paragraph(f"Fecha de emisión: {datetime.now().strftime('%d/%m/%Y')}", subtitulo_style))
    elementos.append(Spacer(1, 10))

    metrics_data = [
        ['Métrica', 'Valor'],
        ['Total de pagos', str(estadisticas['total_pagos'])],
        ['Ingresos totales', f"{estadisticas['ingresos_totales']:.2f}"],
        ['Pago promedio', f"{estadisticas['pago_promedio']:.2f}"],
        ['Pagos anulados', str(estadisticas['pagos_anulados'])],
    ]
    metrics_table = Table(metrics_data, hAlign='LEFT', colWidths=[80 * mm, 65 * mm])
    metrics_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
    ]))
    elementos.extend([metrics_table, Spacer(1, 12)])

    if estadisticas['pagos_por_tipo']:
        elementos.append(Paragraph('Pagos por tipo', styles['Heading3']))
        tipo_data = [['Tipo', 'Pagos', 'Ingresos']]
        for item in estadisticas['pagos_por_tipo']:
            tipo_data.append([item['nombre_tipo_pago'], item['total_pagos'], f"{item['ingresos']:.2f}"])
        tipo_table = Table(tipo_data, hAlign='LEFT', colWidths=[75 * mm, 35 * mm, 45 * mm], repeatRows=1)
        tipo_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
        ]))
        elementos.extend([tipo_table, Spacer(1, 12)])

    if estadisticas['pagos_por_moneda']:
        elementos.append(Paragraph('Pagos por moneda', styles['Heading3']))
        moneda_data = [['Moneda', 'Pagos', 'Ingresos']]
        for item in estadisticas['pagos_por_moneda']:
            moneda_data.append([item['nombre_moneda'], item['total_pagos'], f"{item['ingresos']:.2f}"])
        moneda_table = Table(moneda_data, hAlign='LEFT', colWidths=[75 * mm, 35 * mm, 45 * mm], repeatRows=1)
        moneda_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
        ]))
        elementos.extend([moneda_table, Spacer(1, 12)])

    if estadisticas['top_clientes']:
        elementos.append(Paragraph('Top clientes', styles['Heading3']))
        clientes_data = [['Cliente', 'Pagos', 'Ingresos']]
        for cliente in estadisticas['top_clientes']:
            clientes_data.append([cliente['nombre_cliente'], cliente['total_pagos'], f"{cliente['ingresos']:.2f}"])
        clientes_table = Table(clientes_data, hAlign='LEFT', colWidths=[90 * mm, 35 * mm, 40 * mm], repeatRows=1)
        clientes_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.lightgrey),
            ('GRID', (0, 0), (-1, -1), 0.25, colors.grey),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 2),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.whitesmoke, colors.lightgrey])
        ]))
        elementos.extend([clientes_table, Spacer(1, 12)])

    doc.build(
        elementos,
        onFirstPage=lambda c, d: (borde(c, d), footer(c, d)),
        onLaterPages=lambda c, d: (borde(c, d), footer(c, d))
    )
    return archivo


def _generar_excel(empresa, estadisticas, filtros):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Estadísticas'

    logo_path = _logo_path()
    if os.path.exists(logo_path):
        logo = ExcelImage(logo_path)
        logo.width = 75
        logo.height = 60
        sheet.add_image(logo, 'B1')

    sheet.merge_cells('A6:C6')
    sheet['A6'] = 'ESTUDIO DE BELLEZA GUILLERMINA'
    sheet['A6'].font = Font(bold=True, size=14)
    sheet['A6'].alignment = Alignment(horizontal='center')
    for row, value in enumerate([
        empresa.get('direccion', ''),
        empresa.get('telefono', ''),
        empresa.get('email', ''),
        datetime.now().strftime('%d-%m-%Y'),
    ], start=7):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.cell(row=row, column=1).value = value
        sheet.cell(row=row, column=1).alignment = Alignment(horizontal='center')

    sheet.row_dimensions[1].height = 50
    sheet.row_dimensions[2].height = 8
    sheet.row_dimensions[3].height = 8
    sheet.row_dimensions[4].height = 8
    sheet.row_dimensions[5].height = 8
    sheet.row_dimensions[6].height = 22
    gold_border = Border(bottom=Side(style='medium', color='D4AF37'))
    for cell in sheet[10]:
        cell.border = gold_border
    sheet['A10'].font = Font(italic=True, color='777777')

    sheet.append([])
    sheet.append(['Reporte Estadístico de Pagos'])
    sheet.merge_cells('A12:C12')
    sheet['A12'].font = Font(bold=True, size=14, color='B8860B')
    sheet['A12'].alignment = Alignment(horizontal='center')
    sheet.append(['Emitido', datetime.now().strftime('%Y-%m-%d %H:%M:%S')])
    sheet.append([])
    sheet.append(['Métrica', 'Valor'])
    sheet.append(['Total de pagos', estadisticas['total_pagos']])
    sheet.append(['Ingresos totales', float(estadisticas['ingresos_totales'])])
    sheet.append(['Pago promedio', float(estadisticas['pago_promedio'])])
    sheet.append(['Pagos anulados', estadisticas['pagos_anulados']])
    sheet.append([])

    tipo_sheet = workbook.create_sheet('Pagos por tipo')
    tipo_sheet.append(['Tipo', 'Pagos', 'Ingresos'])
    for item in estadisticas['pagos_por_tipo']:
        tipo_sheet.append([item['nombre_tipo_pago'], item['total_pagos'], float(item['ingresos'])])

    moneda_sheet = workbook.create_sheet('Pagos por moneda')
    moneda_sheet.append(['Moneda', 'Pagos', 'Ingresos'])
    for item in estadisticas['pagos_por_moneda']:
        moneda_sheet.append([item['nombre_moneda'], item['total_pagos'], float(item['ingresos'])])

    top_sheet = workbook.create_sheet('Top clientes')
    top_sheet.append(['Cliente', 'Pagos', 'Ingresos'])
    for item in estadisticas['top_clientes']:
        top_sheet.append([item['nombre_cliente'], item['total_pagos'], float(item['ingresos'])])

    for sheet_to_style in [sheet, tipo_sheet, moneda_sheet, top_sheet]:
        if sheet_to_style.title == 'Estadísticas':
            header_row = 15
        else:
            header_row = 1

        for col in range(1, 4):
            cell = sheet_to_style.cell(row=header_row, column=col)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1B365D')
            cell.alignment = Alignment(horizontal='center')

        sheet_to_style.column_dimensions['A'].width = 24
        sheet_to_style.column_dimensions['B'].width = 14
        sheet_to_style.column_dimensions['C'].width = 16

    archivo = os.path.join(tempfile.gettempdir(), f'reporte_pagos_estadistico_{int(datetime.now().timestamp())}.xlsx')
    workbook.save(archivo)
    return archivo


def handle_request(ruta):
    modelo = PagosModelo()

    desde = request.form.get('desde') or request.args.get('desde')
    hasta = request.form.get('hasta') or request.args.get('hasta')
    cod_tipo_pago = int(request.form.get('cod_tipo_pago', 0) or 0)
    status = request.form.get('status', 'all')

    filtros = {
        'desde': desde or '',
        'hasta': hasta or '',
        'cod_tipo_pago': cod_tipo_pago,
        'status': status,
    }

    estadisticas = _consultar_estadisticas(
        modelo,
        desde=desde,
        hasta=hasta,
        cod_tipo_pago=cod_tipo_pago,
        status=status,
    )
    tipopagos = modelo.listar_tipos_pago()

    if request.method == 'POST' and request.form.get('pdf'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_pdf(empresa_data, estadisticas, filtros)
        return send_file(archivo, mimetype='application/pdf', as_attachment=False)

    if request.method == 'POST' and request.form.get('excel'):
        empresa = EmpresaModelo().mostrar()
        empresa_data = empresa[0] if empresa else {}
        archivo = _generar_excel(empresa_data, estadisticas, filtros)
        return send_file(archivo, as_attachment=True, download_name='reporte_pagos_estadistico.xlsx')

    return {
        'estadisticas': estadisticas,
        'tipopagos': tipopagos,
        'filtros': filtros,
    }
