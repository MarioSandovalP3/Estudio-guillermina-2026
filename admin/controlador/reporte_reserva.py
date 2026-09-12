from flask import request, send_file, session
from admin.modelo.reservas import ReservasModelo
from admin.modelo.empresa import EmpresaModelo

from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Table,
    TableStyle, Spacer, Image
)
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.graphics.shapes import Drawing, Line

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from datetime import datetime
import tempfile
import os


def footer(canvas, doc, empresa):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D4AF37"))
    canvas.setLineWidth(0.8)
    canvas.line(15 * mm, 18 * mm, 195 * mm, 18 * mm)
    canvas.setFont("Helvetica", 8)
    canvas.setFillColor(colors.HexColor("#666666"))
    canvas.drawString(18 * mm, 11 * mm, "Sistema de Reservas • Peluquería")
    canvas.drawRightString(
        195 * mm, 11 * mm,
        datetime.now().strftime("%d/%m/%Y %H:%M")
    )
    canvas.drawCentredString(105 * mm, 6 * mm, f"Página {doc.page}")
    canvas.restoreState()


def borde(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(colors.HexColor("#D4AF37"))
    canvas.setLineWidth(1.2)
    canvas.rect(
        10 * mm, 10 * mm,
        LETTER[0] - 20 * mm,
        LETTER[1] - 20 * mm
    )
    canvas.restoreState()


def handle_request(ruta):
    modelo = ReservasModelo()

    if request.method == "POST" and request.form.get("pdf"):
        empresa_data = EmpresaModelo().mostrar()
        empresa = empresa_data[0] if empresa_data else {}
        listado = modelo.consultar()

        archivo = os.path.join(
            tempfile.gettempdir(),
            "reporte_reservas.pdf"
        )

        doc = SimpleDocTemplate(
            archivo,
            pagesize=LETTER,
            leftMargin=18 * mm,
            rightMargin=18 * mm,
            topMargin=20 * mm,
            bottomMargin=25 * mm
        )

        doc.title = f"Reservas - {empresa.get('nombre', 'Sistema')}"
        doc.author = empresa.get("nombre", "Sistema")
        doc.subject = "Reporte de reservas"

        estilos = getSampleStyleSheet()

        titulo = ParagraphStyle(
            "titulo",
            parent=estilos["Heading1"],
            fontSize=20,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#B8860B"),
            spaceAfter=8
        )

        subtitulo = ParagraphStyle(
            "subtitulo",
            parent=estilos["Normal"],
            fontSize=10,
            alignment=TA_CENTER,
            textColor=colors.grey
        )

        normal = ParagraphStyle(
            "normal",
            parent=estilos["Normal"],
            fontSize=9
        )

        elementos = []
        fecha = datetime.now().strftime("%d/%m/%Y")
        logo = ""

        if empresa.get("logo"):
            base_dir = os.path.dirname(os.path.abspath(__file__))
            logo_bd = empresa.get("logo", "").lstrip("/")
            logo_path = os.path.abspath(
                os.path.join(base_dir, "..", "..", logo_bd)
            )

            if os.path.exists(logo_path):
                logo = Image(
                    logo_path,
                    width=42 * mm,
                    height=42 * mm
                )

        header_text = f"""
        <font size="16" color="#1B365D">
            <b>{empresa.get("nombre", "EMPRESA")}</b>
        </font>
        <br/>
        <font size="9">
            📍 {empresa.get("direccion", "")}<br/>
            📞 {empresa.get("telefono", "")}<br/>
            ✉ {empresa.get("email", "")}
        </font>
        """

        header = Table(
            [[
                logo,
                Paragraph(header_text, normal)
            ]],
            colWidths=[45 * mm, 120 * mm]
        )

        header.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 10)
        ]))

        elementos.append(header)
        elementos.append(Spacer(1, 10))

        line = Drawing(500, 10)
        line.add(Line(
            0, 5, 500, 5,
            strokeColor=colors.HexColor("#D4AF37"),
            strokeWidth=1.2
        ))

        elementos.append(line)
        elementos.append(Spacer(1, 10))

        elementos.append(
            Paragraph("REPORTE GENERAL DE RESERVAS", titulo)
        )

        elementos.append(
            Paragraph("Listado del sistema", subtitulo)
        )

        elementos.append(Spacer(1, 10))

        resumen = [[
            "Fecha",
            fecha,
            "Total",
            str(len(listado))
        ]]

        tabla_resumen = Table(resumen)

        tabla_resumen.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F5F5")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ALIGN", (0, 0), (-1, -1), "CENTER")
        ]))

        elementos.append(tabla_resumen)
        elementos.append(Spacer(1, 12))

        data = [[
            "CLIENTE",
            "SERVICIOS",
            "ESPECIALISTA",
            "FECHA",
            "HORA",
            "ESTADO"
        ]]

        for r in listado:
            estado = (
                "Activo"
                if r["estado_reserva"] == 1
                else "Inactivo"
            )

            data.append([
                Paragraph(
                    f'{r["nombre_cliente"]} {r["apellido_cliente"]}',
                    normal
                ),
                Paragraph(r["servicios"], normal),
                Paragraph(r["nombre_especialista"], normal),
                str(r["fecha_reserva"]),
                str(r["hora_reserva"]),
                estado
            ])

        tabla = Table(data, repeatRows=1)

        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B365D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (0, 0), (-1, -1), "CENTER"),
            ("FONTSIZE", (0, 1), (-1, -1), 8),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#F7F7F7")]
            ),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey)
        ]))

        elementos.append(tabla)

        def first_page(canvas, doc):
            canvas.setTitle(
                f"Reservas - {empresa.get('nombre', 'Sistema')}"
            )
            canvas.setAuthor(
                empresa.get("nombre", "Sistema")
            )
            canvas.setSubject("Reporte de reservas")
            borde(canvas, doc)
            footer(canvas, doc, empresa)

        def later_pages(canvas, doc):
            borde(canvas, doc)
            footer(canvas, doc, empresa)

        doc.build(
            elementos,
            onFirstPage=first_page,
            onLaterPages=later_pages
        )

        return send_file(
            archivo,
            mimetype="application/pdf",
            download_name="reporte_reservas.pdf",
            as_attachment=False
        )

    if request.method == "POST" and request.form.get("excel"):
        empresa_data = EmpresaModelo().mostrar()
        empresa = empresa_data[0] if empresa_data else {}
        listado = modelo.consultar()

        archivo = os.path.join(
            tempfile.gettempdir(),
            "reporte_reservas.xlsx"
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Reservas"

        ws.merge_cells("A1:G1")
        ws["A1"] = empresa.get("nombre", "EMPRESA")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")

        ws["A2"] = empresa.get("direccion", "")
        ws["A3"] = empresa.get("telefono", "")
        ws["A4"] = empresa.get("email", "")
        ws["A5"] = datetime.now().strftime("%d-%m-%Y")

        fila = 7

        headers = [
            "Cédula",
            "Cliente",
            "Servicios",
            "Especialista",
            "Fecha",
            "Hora",
            "Estado"
        ]

        for c, h in enumerate(headers, 1):
            cell = ws.cell(fila, c)
            cell.value = h
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid",
                fgColor="1B365D"
            )
            cell.alignment = Alignment(horizontal="center")

        fila += 1

        for r in listado:
            estado = (
                "Activo"
                if r["estado_reserva"] == 1
                else "Inactivo"
            )

            ws.cell(fila, 1, r["cedula_cliente"])
            ws.cell(
                fila,
                2,
                f'{r["nombre_cliente"]} {r["apellido_cliente"]}'
            )
            ws.cell(fila, 3, r["servicios"])
            ws.cell(fila, 4, r["nombre_especialista"])
            ws.cell(fila, 5, str(r["fecha_reserva"]))
            ws.cell(fila, 6, str(r["hora_reserva"]))
            ws.cell(fila, 7, estado)

            fila += 1

        wb.save(archivo)

        return send_file(
            archivo,
            mimetype=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            download_name="reporte_reservas.xlsx",
            as_attachment=True
        )

    if (
        request.method == "POST"
        and request.form.get("pdf_especialista") == "1"
    ):
        cedula_especialista = request.form.get(
            "especialista_reporte"
        )

        if not cedula_especialista:
            return "Debe seleccionar un especialista", 400

        especialista = modelo.consultar_reporte_especialista(
            cedula_especialista
        )

        if not especialista:
            return "Especialista no encontrado", 404

        empresa = {
            "nombre": session.get("n_empresa", "EMPRESA"),
            "direccion": session.get("direccion_empresa", ""),
            "telefono": session.get("telefono_empresa", ""),
            "email": session.get("email_empresa", ""),
            "logo": session.get("logo", "")
        }

        archivo = os.path.join(
            tempfile.gettempdir(),
            "reporte_especialista.pdf"
        )

        doc = SimpleDocTemplate(
            archivo,
            pagesize=LETTER,
            leftMargin=15 * mm,
            rightMargin=15 * mm,
            topMargin=18 * mm,
            bottomMargin=25 * mm
        )

        estilos = getSampleStyleSheet()

        titulo = ParagraphStyle(
            "titulo_especialista",
            parent=estilos["Heading1"],
            fontSize=18,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#B8860B"),
            spaceAfter=10
        )

        subtitulo = ParagraphStyle(
            "subtitulo_especialista",
            parent=estilos["Normal"],
            fontSize=10,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#666666"),
            spaceAfter=10
        )

        normal = ParagraphStyle(
            "normal_especialista",
            parent=estilos["Normal"],
            fontSize=8
        )

        elementos = []

        logo = ""

        if empresa.get("logo"):
            base_dir = os.path.dirname(
                os.path.abspath(__file__)
            )

            logo_bd = empresa.get(
                "logo",
                ""
            ).lstrip("/")

            logo_path = os.path.abspath(
                os.path.join(
                    base_dir,
                    "..",
                    "..",
                    logo_bd
                )
            )

            if os.path.exists(logo_path):
                logo = Image(
                    logo_path,
                    width=35 * mm,
                    height=35 * mm
                )

        header_text = f"""
        <font size="16" color="#1B365D">
            <b>{empresa.get("nombre", "EMPRESA")}</b>
        </font>
        <br/>
        <font size="9">
            📍 {empresa.get("direccion", "")}<br/>
            📞 {empresa.get("telefono", "")}<br/>
            ✉ {empresa.get("email", "")}
        </font>
        """

        header = Table(
            [[
                logo,
                Paragraph(header_text, normal)
            ]],
            colWidths=[45 * mm, 120 * mm]
        )

        header.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 8)
        ]))

        elementos.append(header)
        elementos.append(Spacer(1, 5))

        linea = Drawing(500, 10)

        linea.add(Line(
            0,
            5,
            500,
            5,
            strokeColor=colors.HexColor("#D4AF37"),
            strokeWidth=1.2
        ))

        elementos.append(linea)
        elementos.append(Spacer(1, 10))

        elementos.append(
            Paragraph(
                "REPORTE DEL ESPECIALISTA",
                titulo
            )
        )

        elementos.append(
            Paragraph(
                f'{especialista.get("especialista", "")} — '
                f'{especialista.get("nombre_usuario", "")} '
                f'{especialista.get("apellido_usuario", "")}',
                subtitulo
            )
        )

        elementos.append(Spacer(1, 8))

        datos_especialista = [
            [
                "Especialista",
                (
                    f'{especialista.get("especialista", "")} — '
                    f'{especialista.get("nombre_usuario", "")} '
                    f'{especialista.get("apellido_usuario", "")}'
                )
            ],
            [
                "Cédula",
                str(
                    especialista.get(
                        "cedula_especialista",
                        ""
                    )
                )
            ],
            [
                "Total de reservas",
                str(
                    especialista.get(
                        "total_reservas",
                        0
                    )
                )
            ],
            [
                "Cliente más atendido",
                especialista.get(
                    "cliente_mas_atendido",
                    "No disponible"
                )
            ],
            [
                "Servicio más solicitado",
                especialista.get(
                    "servicio_mas_solicitado",
                    "No disponible"
                )
            ]
        ]

        tabla_info = Table(
            datos_especialista,
            colWidths=[
                55 * mm,
                115 * mm
            ]
        )

        tabla_info.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#1B365D")),
            ("TEXTCOLOR", (0, 0), (0, -1), colors.white),
            ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.lightgrey),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5)
        ]))

        elementos.append(tabla_info)
        elementos.append(Spacer(1, 15))

        elementos.append(
            Paragraph(
                "RESERVAS REALIZADAS",
                ParagraphStyle(
                    "titulo_reservas",
                    parent=estilos["Heading2"],
                    fontSize=12,
                    textColor=colors.HexColor("#1B365D"),
                    spaceAfter=8
                )
            )
        )

        data = [[
            "CLIENTE",
            "FECHA",
            "HORA",
            "PROMOCIÓN",
            "SERVICIOS",
            "TOTAL"
        ]]

        reservas = especialista.get("reservas", [])

        for reserva in reservas:
            servicios_precios = reserva.get(
                "servicios_precios",
                ""
            )

            servicios_formateados = []

            if servicios_precios:
                servicios_lista = servicios_precios.split("@@")

                for servicio in servicios_lista:
                    partes = servicio.split("###")

                    if len(partes) >= 2:
                        nombre = partes[0]
                        precio = partes[1]

                        servicios_formateados.append(
                            f'{nombre} '
                            f'<font color="#777777">'
                            f'${precio}'
                            f'</font>'
                        )

            servicios_html = (
                "<br/>".join(servicios_formateados)
                if servicios_formateados
                else "No disponible"
            )

            promocion = reserva.get(
                "promocion",
                "NO"
            )

            if not promocion:
                promocion = "NO"

            total = reserva.get("total", 0)

            try:
                total_formateado = f"${float(total):,.2f}"
            except Exception:
                total_formateado = f"${total}"

            data.append([
                Paragraph(
                    reserva.get(
                        "cliente",
                        "No disponible"
                    ),
                    normal
                ),
                str(
                    reserva.get(
                        "fecha",
                        ""
                    )
                ),
                str(
                    reserva.get(
                        "hora",
                        ""
                    )
                ),
                Paragraph(
                    str(promocion),
                    normal
                ),
                Paragraph(
                    servicios_html,
                    normal
                ),
                total_formateado
            ])

        if len(data) == 1:
            data.append([
                Paragraph(
                    "No existen reservas",
                    normal
                ),
                "",
                "",
                "",
                "",
                "$0.00"
            ])

        tabla_reservas = Table(
            data,
            repeatRows=1,
            colWidths=[
                36 * mm,
                21 * mm,
                17 * mm,
                29 * mm,
                65 * mm,
                20 * mm
            ]
        )

        tabla_reservas.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1B365D")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, 0), 7),
            ("FONTSIZE", (0, 1), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (1, 1), (3, -1), "CENTER"),
            ("ALIGN", (5, 1), (5, -1), "RIGHT"),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [colors.white, colors.HexColor("#F7F7F7")]
            ),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4)
        ]))

        elementos.append(tabla_reservas)

        def first_page(canvas, doc):
            canvas.setTitle("Reporte del especialista")
            canvas.setAuthor(
                empresa.get("nombre", "Sistema")
            )
            canvas.setSubject(
                "Reporte individual del especialista"
            )
            borde(canvas, doc)
            footer(canvas, doc, empresa)

        def later_pages(canvas, doc):
            borde(canvas, doc)
            footer(canvas, doc, empresa)

        doc.build(
            elementos,
            onFirstPage=first_page,
            onLaterPages=later_pages
        )

        return send_file(
            archivo,
            mimetype="application/pdf",
            download_name="reporte_especialista.pdf",
            as_attachment=False
        )

    reservas = modelo.consultar()
    especialistas = modelo.consultar_especialistas()

    return {
        "reservas": reservas,
        "especialistas": especialistas
    }